"""Build the SC4020 Project 1 report from the results in results/*.json.

Every number in the report (tables and prose) is read from the results JSONs
at build time, and method parameters are imported from the method files, so
the report cannot drift from the code. Outputs:
    reports/figures/*.png          report versions of the figures
    reports/group_XX_report.docx   editable source
    reports/group_XX_report.pdf    via Microsoft Word (with --pdf)

Run from the repo root:
    venv/bin/python reports/build_report.py --pdf

Needs python-docx, docx2pdf and pypdf (report tooling only, not in
requirements.txt). Re-running overwrites the .docx, so make wording changes
here, or hand-edit the .docx only after the final build.
"""

import argparse
import importlib.util
import json
import os
import re
import sys

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
REPORT_DIR = os.path.join(ROOT, "reports")
FIG_DIR = os.path.join(REPORT_DIR, "figures")
RESULTS = os.path.join(ROOT, "results")

# Shared loader, Pareto helper and palette from the combined-plot script.
_spec = importlib.util.spec_from_file_location("pc", os.path.join(RESULTS, "plot_combined.py"))
pc = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pc)

# Method parameters, imported so the report always states what the code ran.
import methods.hnsw as hnsw_mod  # noqa: E402
import methods.annoy_method as annoy_mod  # noqa: E402
import methods.lsh as lsh_mod  # noqa: E402
import eval.run_PQ_IVF as pqivf_mod  # noqa: E402

DS = ["sift1m", "wikipedia"]
DS_NAME = {"sift1m": "SIFT1M", "wikipedia": "Wikipedia"}
METHODS = ["Linear scan", "LSH", "PQ-ADC", "PQ-SDC", "IVF", "IVF+PQ", "HNSW", "Annoy"]


# --------------------------------------------------------------------------- #
# Data
# --------------------------------------------------------------------------- #

ROWS = pc.load_rows()
with open(os.path.join(RESULTS, "lsh_results.json"), encoding="utf-8") as f:
    LSH_RAW = {ds: v["results"] for ds, v in json.load(f).items()}
with open(os.path.join(RESULTS, "annoy_sift1m_metric_comparison.json"), encoding="utf-8") as f:
    ABLATION = json.load(f)
with open(os.path.join(RESULTS, "annoy_default_search_k.json"), encoding="utf-8") as f:
    ANNOY_DEFAULT = {ds: next(iter(v.values())) for ds, v in json.load(f).items()}
with open(os.path.join(RESULTS, "kmeans_seed_sensitivity.json"), encoding="utf-8") as f:
    KMEANS = json.load(f)
with open(os.path.join(RESULTS, "wiki_query_cases.json"), encoding="utf-8") as f:
    CASES = json.load(f)
with open(os.path.join(RESULTS, "timing_controlled.json"), encoding="utf-8") as f:
    RETIME = json.load(f)["rounds"]
with open(os.path.join(RESULTS, "timing_controlled_sift1m.json"), encoding="utf-8") as f:
    RETIME_S = json.load(f)["rounds"]
with open(os.path.join(RESULTS, "sift_cosine_gt_overlap.json"), encoding="utf-8") as f:
    COSINE_GT = json.load(f)

# The superseded 4-hash Wikipedia LSH configs, read from git history (the
# commit before the 8-hash retune) so the report can cite them exactly.
OLD_LSH_COMMIT = "8ec046a"
import subprocess  # noqa: E402
OLD_LSH_W = list(json.loads(subprocess.run(
    ["git", "show", f"{OLD_LSH_COMMIT}:results/lsh_results.json"],
    cwd=ROOT, capture_output=True, text=True, check=True).stdout)["wikipedia"]["results"].values())


def pts(method, ds):
    return [r for r in ROWS if r["method"] == method and r["dataset"] == ds]


def one(method, ds, param):
    return next(r for r in pts(method, ds) if r["param"] == param)


def linear(ds):
    return one("Linear scan", ds, "exact")


def fastest_at(method, ds, threshold=0.9):
    ok = [r for r in pts(method, ds) if r["recall_at_10"] >= threshold]
    return min(ok, key=lambda r: r["avg_query_time_ms"]) if ok else None


def best_recall(method, ds):
    return max(pts(method, ds), key=lambda r: r["recall_at_10"])


def ivf_nprobe(param):
    return int(re.search(r"nprobe(\d+)", param).group(1))


def pq_m(param):
    return int(re.search(r"(?:^|_)m(\d+)", param).group(1))


def dataset_shapes():
    """Read dataset shapes from the data files (headers / mmap only)."""
    sift = os.path.join(ROOT, "data", "sift")

    def fvecs_shape(name, itemsize=4):
        path = os.path.join(sift, name)
        dim = int(np.fromfile(path, dtype=np.int32, count=1)[0])
        return os.path.getsize(path) // (itemsize * (dim + 1)), dim

    wiki = os.path.join(ROOT, "data", "wiki")
    load = lambda n: np.load(os.path.join(wiki, n), mmap_mode="r").shape
    with open(os.path.join(wiki, "wiki_sentences.json"), encoding="utf-8") as f:
        sentences = json.load(f)
    wiki_cfg = sentences["config"]
    # Sentences shared between splits (checked so the report can state the
    # splits are disjoint rather than assume it).
    splits = {k: set(sentences[k]) for k in ("base", "query", "train")}
    wiki_cfg["overlap"] = (len(splits["query"] & splits["base"]) + len(splits["train"] & splits["base"])
                           + len(splits["query"] & splits["train"]))
    return {
        "sift1m": {"base": fvecs_shape("sift_base.fvecs"), "query": fvecs_shape("sift_query.fvecs"),
                   "train": fvecs_shape("sift_learn.fvecs"), "gt": fvecs_shape("sift_groundtruth.ivecs")},
        "wikipedia": {"base": load("wiki_base_embeddings.npy"), "query": load("wiki_query_embeddings.npy"),
                      "train": load("wiki_train_embeddings.npy")},
        "wiki_cfg": wiki_cfg,
    }


SHAPES = dataset_shapes()


# --------------------------------------------------------------------------- #
# Number formatting (one formatter per quantity, used by tables and prose)
# --------------------------------------------------------------------------- #

def fr(x):
    return f"{x:.3f}"


def fms(x):
    return f"{x:.3f}" if x < 1 else f"{x:.2f}" if x < 10 else f"{x:.1f}"


def fmb(x):
    return f"{x:.0f}" if x >= 100 else f"{x:.1f}" if x >= 10 else f"{x:.2f}"


def fs(x):
    return f"{x:.1f}"


def fx(x):
    return f"{x:.0f}×" if x >= 10 else f"{x:.1f}×"


NUM = "zero one two three four five six seven eight nine ten".split()


def fint(x):
    return f"{round(float(x)):,}"


def pts_diff(a, b):
    return f"{100 * (a - b):.1f}"


def pretty(method, ds, param):
    if method in ("HNSW",):
        return f"ef_search={param}"
    if method == "Annoy":
        return f"search_k={param}"
    if method == "IVF":
        return f"nprobe={ivf_nprobe(param)}"
    if method == "IVF+PQ":
        return f"m={pq_m(param)}, nprobe={ivf_nprobe(param)}"
    if method.startswith("PQ"):
        return f"m={pq_m(param)}"
    if method == "LSH":
        c = LSH_RAW[ds][param]
        return f"{param} (L={c['num_tables']}, K={c['num_hashes']}, w={c['bucket_width']:g})"
    return "exact"


def join_and(items):
    items = list(items)
    return items[0] if len(items) == 1 else f"{', '.join(items[:-1])} and {items[-1]}"


def environment():
    """Machine and library versions of the venv the benchmarks ran in."""
    import importlib.metadata as md
    import platform
    import subprocess
    sysctl = lambda key: subprocess.run(["sysctl", "-n", key], capture_output=True, text=True).stdout.strip()
    return {
        "cpu": sysctl("machdep.cpu.brand_string"), "cores": sysctl("hw.ncpu"),
        "mem_gb": int(sysctl("hw.memsize")) // 2**30, "macos": platform.mac_ver()[0],
        "python": platform.python_version(),
        **{p: md.version(p) for p in ("numpy", "faiss-cpu", "hnswlib", "annoy")},
    }


def rng(values, fmt):
    lo, hi = min(values), max(values)
    return fmt(lo) if fmt(lo) == fmt(hi) else f"{fmt(lo)}–{fmt(hi)}"


# --------------------------------------------------------------------------- #
# Claim checks: every claim from the brief is tested against the data
# --------------------------------------------------------------------------- #

CLAIMS = []


def claim(text, ok, detail):
    CLAIMS.append((text, bool(ok), detail))


def check_claims():
    acc_win, acc_loss, fast_win, fast_loss, no_match = [], [], [], [], []
    for ds in DS:
        lin = linear(ds)["avg_query_time_ms"]
        h, v = fastest_at("HNSW", ds), fastest_at("IVF", ds)
        others = [fastest_at(m, ds) for m in METHODS if m not in ("HNSW", "Linear scan")]
        claim(f"HNSW is fastest to recall>=0.9 on {ds}",
              all(o is None or o["avg_query_time_ms"] >= h["avg_query_time_ms"] for o in others),
              f"HNSW {fms(h['avg_query_time_ms'])} ms vs IVF {fms(v['avg_query_time_ms'])} ms")
        for m_ in sorted({pq_m(r["param"]) for r in pts("PQ-ADC", ds)}):
            pq = next(r for r in pts("PQ-ADC", ds) if pq_m(r["param"]) == m_)
            sdc = next(r for r in pts("PQ-SDC", ds) if pq_m(r["param"]) == m_)
            ivfpq = [r for r in pts("IVF+PQ", ds) if pq_m(r["param"]) == m_]
            best = max(ivfpq, key=lambda r: r["recall_at_10"])
            match = [r for r in ivfpq if r["recall_at_10"] >= pq["recall_at_10"]]
            (acc_win if best["recall_at_10"] > pq["recall_at_10"] else acc_loss).append(
                f"{ds} m={m_} {fr(best['recall_at_10'])} vs {fr(pq['recall_at_10'])}")
            if match:  # the report claims speed only where IVF+PQ reaches PQ's recall
                (fast_win if min(r["avg_query_time_ms"] for r in match) < pq["avg_query_time_ms"]
                 else fast_loss).append(f"{ds} m={m_}")
            else:
                no_match.append(f"{ds} m={m_}")
            claim(f"PQ-SDC slower than ADC ({ds}, m={m_})",
                  sdc["avg_query_time_ms"] > pq["avg_query_time_ms"],
                  f"{fms(sdc['avg_query_time_ms'])} vs {fms(pq['avg_query_time_ms'])} ms")
        claim(f"IVF+PQ ceiling ({ds})", True, fr(best_recall("IVF+PQ", ds)["recall_at_10"]))
        claim(f"IVF+PQ index far smaller than IVF ({ds})",
              max(r["index_size_mb"] for r in pts("IVF+PQ", ds)) < pts("IVF", ds)[0]["index_size_mb"] / 5,
              f"{rng([r['index_size_mb'] for r in pts('IVF+PQ', ds)], fmb)} vs {fmb(pts('IVF', ds)[0]['index_size_mb'])} MB")
    claim("IVF+PQ more accurate than PQ-ADC in most code-size settings", len(acc_win) > len(acc_loss),
          f"{len(acc_win)} of {len(acc_win) + len(acc_loss)}; exceptions: {', '.join(acc_loss) or 'none'}")
    claim("IVF+PQ faster than PQ-ADC wherever it matches PQ's recall", fast_win and not fast_loss,
          f"{len(fast_win)} settings; never matches: {', '.join(no_match) or 'none'}")
    lin_w = linear("wikipedia")["avg_query_time_ms"]
    claim("LSH never beats linear scan on Wikipedia",
          min(r["avg_query_time_ms"] for r in pts("LSH", "wikipedia")) > lin_w,
          f"fastest LSH {fms(min(r['avg_query_time_ms'] for r in pts('LSH', 'wikipedia')))} vs {fms(lin_w)} ms")
    a = one("Annoy", "wikipedia", str(max(annoy_mod.SEARCH_K_SWEEP)))
    claim("Near-exact Annoy on Wikipedia not faster than brute force (main run)",
          a["avg_query_time_ms"] > lin_w, f"{fms(a['avg_query_time_ms'])} vs {fms(lin_w)} ms")
    diffs = [100 * (ABLATION["euclidean"][k]["recall_at_10"] - ABLATION["angular"][k]["recall_at_10"])
             for k in ABLATION["euclidean"]]
    claim("Ablation: metric explains at most 1.5 recall points",
          round(max(abs(d) for d in diffs), 1) <= 1.5, f"differences {', '.join(f'{d:.2f}' for d in diffs)} points")
    for ds in DS:
        for method, key in (("Annoy", int), ("HNSW", int), ("IVF", ivf_nprobe)):
            seq = sorted(pts(method, ds), key=lambda r: key(r["param"]))
            mono = all(b["recall_at_10"] >= a_["recall_at_10"] and b["avg_query_time_ms"] >= a_["avg_query_time_ms"]
                       for a_, b in zip(seq, seq[1:]))
            claim(f"{method} recall and time rise with its search parameter ({ds})", mono, "")


# --------------------------------------------------------------------------- #
# Figures (report copies; results/ originals are not touched)
# --------------------------------------------------------------------------- #

def fig_recall_latency(ds, path):
    fig, ax = plt.subplots(figsize=(10.5, 5.6), facecolor="#fcfcfb")
    ax.set_facecolor("#fcfcfb")
    pc.style_axes(ax)
    lin = linear(ds)
    ax.axvline(lin["avg_query_time_ms"], color=pc.MUTED, linestyle="--", linewidth=1, alpha=0.45, zorder=1)
    ax.text(lin["avg_query_time_ms"] * 0.93, 0.02, "linear scan query time", rotation=90,
            ha="right", va="bottom", fontsize=8, color=pc.MUTED)
    for method, (color, marker) in pc.STYLE.items():
        p = pts(method, ds)
        front = pc.pareto(p)
        others = [q for q in p if q not in front]
        ax.scatter([q["avg_query_time_ms"] for q in others], [q["recall_at_10"] for q in others],
                   s=18, marker=marker, color=color, alpha=0.25, linewidths=0)
        ax.plot([q["avg_query_time_ms"] for q in front], [q["recall_at_10"] for q in front],
                color=color, marker=marker, markersize=6, linewidth=2, label=method,
                markeredgecolor="#fcfcfb", markeredgewidth=1)
    ax.scatter(lin["avg_query_time_ms"], lin["recall_at_10"], s=140, marker="*", color=pc.INK,
               zorder=5, label="Linear scan (exact)")
    ax.set_xscale("log")
    ax.set_ylim(0, 1.03)
    ax.set_xlabel("Average query time (ms, log scale)", color=pc.INK)
    ax.set_ylabel("Recall@10", color=pc.INK)
    ax.set_title(f"Recall@10 vs query time: {DS_NAME[ds]} (single-threaded)",
                 loc="left", color=pc.INK, fontsize=12, fontweight="bold", pad=22)
    ax.text(0, 1.01, "Lines: best configs per method (Pareto frontier). Faint markers: dominated configs.",
            transform=ax.transAxes, fontsize=8.5, color=pc.MUTED, va="bottom")
    # Outside the axes: every inside corner collides with some curve on one dataset.
    ax.legend(loc="center left", bbox_to_anchor=(1.01, 0.5), fontsize=8.5, frameon=False, labelcolor=pc.INK)
    fig.tight_layout()
    fig.savefig(path, dpi=300, metadata={"Software": None})
    plt.close(fig)


def fig_index_size(path):
    methods = ["Linear scan", "LSH", "PQ", "IVF", "IVF+PQ", "HNSW", "Annoy"]
    colors = {"sift1m": "#2a78d6", "wikipedia": "#eb6834"}
    fig, ax = plt.subplots(figsize=(9, 5.5), facecolor="#fcfcfb")
    ax.set_facecolor("#fcfcfb")
    pc.style_axes(ax)
    ax.grid(False, axis="x")
    width = 0.38
    for i, ds in enumerate(DS):
        for j, method in enumerate(methods):
            sizes = [r["index_size_mb"] for r in ROWS if r["dataset"] == ds and r["method"].split("-")[0] == method]
            x = j + (i - 0.5) * (width + 0.02)
            lo, hi = min(sizes), max(sizes)
            ax.bar(x, lo, width, color=colors[ds], label=DS_NAME[ds] if j == 0 else None)
            label, top = fmb(lo), lo
            if hi > lo * 1.05:
                ax.plot([x, x], [lo, hi], color=pc.INK, linewidth=1.2)
                ax.plot([x - 0.08, x + 0.08], [hi, hi], color=pc.INK, linewidth=1.2)
                label, top = f"{fmb(lo)}–{fmb(hi)}", hi
            ax.annotate(label, (x, top), xytext=(0, 3), textcoords="offset points",
                        ha="center", fontsize=7.5, color=pc.MUTED)
    ax.set_yscale("log")
    ax.set_xticks(range(len(methods)), methods, color=pc.INK)
    ax.set_ylabel("Index size (MB, log scale)", color=pc.INK)
    ax.set_title("Index size by method", loc="left", color=pc.INK, fontsize=12, fontweight="bold", pad=22)
    ax.text(0, 1.01, "Bar: smallest config. Whisker: largest config, where configs differ.",
            transform=ax.transAxes, fontsize=8.5, color=pc.MUTED, va="bottom")
    ax.legend(frameon=False, fontsize=9, labelcolor=pc.INK, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.07))
    fig.tight_layout()
    fig.savefig(path, dpi=300, metadata={"Software": None})
    plt.close(fig)


def fig_ablation(path):
    ks = sorted(int(k) for k in ABLATION["euclidean"])
    series = [
        ("SIFT1M, euclidean (ablation run)", [ABLATION["euclidean"][str(k)]["recall_at_10"] for k in ks], "#2a78d6", "o"),
        ("SIFT1M, angular (ablation run)", [ABLATION["angular"][str(k)]["recall_at_10"] for k in ks], "#eb6834", "s"),
        ("Wikipedia, angular (final run)", [one("Annoy", "wikipedia", str(k))["recall_at_10"] for k in ks], "#1baf7a", "D"),
    ]
    fig, ax = plt.subplots(figsize=(8, 5), facecolor="#fcfcfb")
    ax.set_facecolor("#fcfcfb")
    pc.style_axes(ax)
    for label, ys, color, marker in series:
        ax.plot(ks, ys, color=color, marker=marker, markersize=6, linewidth=2, label=label,
                markeredgecolor="#fcfcfb", markeredgewidth=1)
    ax.set_xscale("log")
    ax.set_xticks(ks, [fint(k) for k in ks])
    ax.minorticks_off()
    ax.set_ylim(0.5, 1.01)
    ax.set_xlabel("search_k (nodes inspected per query, log scale)", color=pc.INK)
    ax.set_ylabel("Recall@10", color=pc.INK)
    ax.set_title("Annoy: metric choice vs dataset", loc="left", color=pc.INK, fontsize=12, fontweight="bold", pad=22)
    ax.text(0, 1.01, f"n_trees = {annoy_mod.N_TREES}. Recall only: the ablation run used different thread settings.",
            transform=ax.transAxes, fontsize=8.5, color=pc.MUTED, va="bottom")
    ax.legend(loc="lower right", fontsize=9, frameon=False, labelcolor=pc.INK)
    fig.tight_layout()
    fig.savefig(path, dpi=300, metadata={"Software": None})
    plt.close(fig)


def fig_cases(path):
    """Per-query recall on Wikipedia: distribution, and by neighbourhood tightness."""
    st = CASES["settings"]
    series = [("IVF", f"IVF (nprobe {st['ivf_nprobe']})", "recall_ivf", "mean_recall_ivf"),
              ("IVF+PQ", f"IVF+PQ (m {st['ivfpq_m']}, nprobe {st['ivfpq_nprobe']})", "recall_ivfpq", "mean_recall_ivfpq")]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4.4), facecolor="#fcfcfb")
    levels = np.round(np.arange(0, 1.01, 0.1), 1)
    width = 0.042
    for i, (method, label, key, _) in enumerate(series):
        rec = np.round(np.array(CASES["per_query"][key]), 1)
        counts = [int((rec == lv).sum()) for lv in levels]
        a1.bar(levels + (i - 0.5) * (width + 0.004), counts, width, color=pc.STYLE[method][0], label=label)
    a1.set_xticks(levels, [f"{lv:.1f}" for lv in levels])
    a1.set_xlabel("Per-query recall@10", color=pc.INK)
    a1.set_ylabel("Number of queries", color=pc.INK)
    a1.set_title("(a) Distribution of per-query recall", loc="left", color=pc.INK, fontsize=11, fontweight="bold")
    a1.legend(frameon=False, fontsize=8.5, labelcolor=pc.INK, loc="upper left")
    qs = CASES["s10_quartiles"]
    x = np.arange(len(qs))
    for method, label, _, mkey in series:
        color, marker = pc.STYLE[method]
        a2.plot(x, [q[mkey] for q in qs], color=color, marker=marker, markersize=7, linewidth=2, label=label,
                markeredgecolor="#fcfcfb", markeredgewidth=1)
    a2.set_xticks(x, [f"Q{i + 1}: {q['s10_range'][0]:.2f}–{q['s10_range'][1]:.2f}\n{q['mean_n_cells']:.1f} cells"
                      for i, q in enumerate(qs)], fontsize=8)
    a2.set_ylim(0.6, 1.0)
    a2.set_xlabel("Quartile of 10th-neighbour cosine similarity (loose → tight)\nmean IVF cells holding the true top-10",
                  color=pc.INK, fontsize=9)
    a2.set_ylabel("Mean recall@10", color=pc.INK)
    a2.set_title("(b) Recall vs neighbourhood tightness", loc="left", color=pc.INK, fontsize=11, fontweight="bold")
    a2.legend(frameon=False, fontsize=8.5, labelcolor=pc.INK, loc="lower right")
    for ax in (a1, a2):
        ax.set_facecolor("#fcfcfb")
        pc.style_axes(ax)
    a1.grid(False, axis="x")
    fig.tight_layout()
    fig.savefig(path, dpi=300, metadata={"Software": None})
    plt.close(fig)


# --------------------------------------------------------------------------- #
# DOCX helpers
# --------------------------------------------------------------------------- #

FONT = "Times New Roman"
FIG = {"sift": 1, "wiki": 2, "size": 3, "ablation": 4, "cases": 5}
TAB = {"methods": 1, "data": 2, "summary": 3, "sweeps": 4, "lsh": 5, "pq": 6, "kmeans": 7, "ivfpq": 8,
       "ablation": 9, "cases": 10, "strengths": 11}


def F(key):
    return f"Figure {FIG[key]}"


def T(key):
    return f"Table {TAB[key]}"


TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\[CHECK[^\]]*\])")


def add_runs(par, text, size=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1]); run.font.name = "Courier New"
            run.font.size = Pt((size or 11) - 1.5)
            continue
        elif part.startswith("[CHECK"):
            run = par.add_run(part); run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("*"):
            run = par.add_run(part[1:-1]); run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)


def set_cell_borders(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, size in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if size else "nil")
        if size:
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), "000000")
        borders.append(el)


class Report:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.5))
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size in (("Heading 1", 13.5), ("Heading 2", 11.5), ("Title", 17)):
            st = self.doc.styles[name]
            rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(attr), FONT)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rfonts.get(qn(attr)) is not None:
                    del rfonts.attrib[qn(attr)]
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.paragraph_format.space_before = Pt(12 if name != "Heading 2" else 8)
            st.paragraph_format.space_after = Pt(4)
        # Page numbers in the footer.
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; r.append(t); fld.append(r)
        fp._p.append(fld)

    def title(self, text, lines):
        p = self.doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, text)
        for line in lines:
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(2)
            add_runs(q, line)

    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def p(self, text, size=None, align_justify=True):
        par = self.doc.add_paragraph()
        if align_justify:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(par, text, size)
        return par

    def bullets(self, items):
        for item in items:
            par = self.doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            add_runs(par, item)

    def caption(self, kind, key, text):
        num = FIG[key] if kind == "Figure" else TAB[key]
        par = self.doc.add_paragraph()
        par.paragraph_format.space_after = Pt(10 if kind == "Figure" else 4)
        par.paragraph_format.keep_together = True  # never split a caption across pages
        if kind == "Table":
            par.paragraph_format.keep_with_next = True
        run = par.add_run(f"{kind} {num}. "); run.bold = True; run.font.size = Pt(9.5)
        add_runs(par, text, 9.5)

    def figure(self, key, path, caption, width_cm=15.5):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        # Single spacing: Word scales an inline image's line by the 1.15 body spacing.
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.caption("Figure", key, caption)

    def table(self, key, caption, header, rows, align=None, note=None, widths=None, split_ok=False):
        self.caption("Table", key, caption)
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        align = align or ["l"] + ["r"] * (len(header) - 1)
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
        for i, row in enumerate([header] + rows):
            is_group = isinstance(row, str)
            cells = tbl.rows[i].cells
            # Rows never split and the header repeats on a new page. Short tables
            # stay on one page (each row kept with the next); long ones (split_ok)
            # may break between rows, but the header and group labels stay with
            # the row below them.
            tr_pr = tbl.rows[i]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                tr_pr.append(OxmlElement("w:tblHeader"))
            # split_ok: keep each group block together, breaking only before a group label.
            keep = i < len(rows) and (not split_ok or i == 0 or is_group or not isinstance(rows[i], str))
            if is_group:  # full-width group label row
                merged = cells[0].merge(cells[-1])
                merged.text = ""
                par = merged.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.keep_with_next = keep
                run = par.add_run(row); run.bold = True; run.italic = True; run.font.size = Pt(9)
                set_cell_borders(merged, top=4)
                continue
            for j, value in enumerate(row):
                cell = cells[j]
                cell.text = ""
                par = cell.paragraphs[0]
                par.alignment = amap[align[j]]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.keep_with_next = keep
                add_runs(par, str(value), 9)
                if i == 0:
                    for run in par.runs:
                        run.bold = True
        for cell in tbl.rows[0].cells:
            set_cell_borders(cell, top=12, bottom=6)
        for cell in tbl.rows[-1].cells:
            set_cell_borders(cell, bottom=12)
        if widths:
            tbl.autofit = False
            for row in tbl.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
        if note:
            par = self.doc.add_paragraph()
            par.paragraph_format.space_before = Pt(3)
            add_runs(par, note, 8.5)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def save(self, path):
        self.doc.save(path)


# --------------------------------------------------------------------------- #
# Report content
# --------------------------------------------------------------------------- #

def build(path_docx):
    S, W = "sift1m", "wikipedia"
    lin = {ds: linear(ds) for ds in DS}
    lin_ms = {ds: lin[ds]["avg_query_time_ms"] for ds in DS}
    fast = {(m, ds): fastest_at(m, ds) for m in METHODS for ds in DS}
    hn = {ds: fast[("HNSW", ds)] for ds in DS}
    iv = {ds: fast[("IVF", ds)] for ds in DS}
    an = {ds: fast[("Annoy", ds)] for ds in DS}
    lshw = fast[("LSH", W)]
    ivfpq_best = {ds: best_recall("IVF+PQ", ds) for ds in DS}
    pq_sizes = {ds: [r["index_size_mb"] for r in pts("PQ-ADC", ds)] for ds in DS}
    ivfpq_sizes = {ds: [r["index_size_mb"] for r in pts("IVF+PQ", ds)] for ds in DS}
    ivf_size = {ds: pts("IVF", ds)[0]["index_size_mb"] for ds in DS}
    hnsw_build = {ds: pts("HNSW", ds)[0]["build_time_sec"] for ds in DS}
    ivf_build = {ds: pts("IVF", ds)[0]["build_time_sec"] for ds in DS}
    annoy_build = {ds: pts("Annoy", ds)[0]["build_time_sec"] for ds in DS}
    sh = SHAPES
    nq = {ds: sh[ds]["query"][0] for ds in DS}
    nb = {ds: sh[ds]["base"][0] for ds in DS}
    wcfg = sh["wiki_cfg"]
    lsh_w = {n: LSH_RAW[W][n] for n in LSH_RAW[W]}
    lsh_s = {n: LSH_RAW[S][n] for n in LSH_RAW[S]}
    lsh_w_low = lsh_w["low"]
    lsh_s_low = lsh_s["low"]
    ns_per_cand_w = lsh_w_low["avg_query_time_ms"] * 1e6 / lsh_w_low["avg_candidates"]
    ns_per_vec_w = lin_ms[W] * 1e6 / nb[W]
    ns_per_cand_s = lsh_s_low["avg_query_time_ms"] * 1e6 / lsh_s_low["avg_candidates"]
    ns_per_vec_s = lin_ms[S] * 1e6 / nb[S]
    annoy_max = {ds: one("Annoy", ds, str(max(annoy_mod.SEARCH_K_SWEEP))) for ds in DS}
    annoy_min = {ds: one("Annoy", ds, str(min(annoy_mod.SEARCH_K_SWEEP))) for ds in DS}

    # Controlled Wikipedia re-timings (results/timing_controlled.json): rt holds
    # each configuration's per-round times; rt_sp the speed-up over linear scan
    # within each round, the comparison that survives machine-wide slowdowns.
    rt = {n: np.array([r["avg_query_time_ms"][n] for r in RETIME]) for n in RETIME[0]["avg_query_time_ms"]}
    rt_lin = rt["linear_scan"]
    rt_sp = {n: rt_lin / t for n, t in rt.items()}
    rt_spread = {n: 100 * (t.max() - t.min()) / np.median(t) for n, t in rt.items()}
    n_rt = len(rt_lin)
    rt_load = [x for r in RETIME for x in r["loadavg"]]
    rt_main = {"linear_scan": lin_ms[W], "lsh_low": lsh_w_low["avg_query_time_ms"],
               "ivf_nprobe16": iv[W]["avg_query_time_ms"], "hnsw_ef64": hn[W]["avg_query_time_ms"],
               "annoy_search_k20000": one("Annoy", W, "20000")["avg_query_time_ms"],
               "annoy_search_k100000": annoy_max[W]["avg_query_time_ms"]}
    rt_main_faster = [n for n in rt if rt_main[n] < rt[n].min()]
    rt_fast = ("ivf_nprobe16", "hnsw_ef64")  # the sub-millisecond configurations
    rt_slow = [n for n in rt if n not in rt_fast]
    rt_short_s = [np.median(rt[n]) for n in rt_fast]  # ms/query x 1000 queries = s

    claim("Re-timings: sub-ms configs vary more than the rest",
          min(rt_spread[n] for n in rt_fast) > max(rt_spread[n] for n in rt_slow),
          ", ".join(f"{n} {v:.0f}%" for n, v in rt_spread.items()))
    claim("Re-timings: LSH and near-exact Annoy level with linear scan every round",
          all(0.8 < x < 1.25 for n in ("lsh_low", "annoy_search_k100000") for x in rt_sp[n]),
          f"LSH {rt_sp['lsh_low'].round(2)}, Annoy {rt_sp['annoy_search_k100000'].round(2)}")

    def rt_range(n, prec=0, sp_=None):
        sp = (sp_ or rt_sp)[n]
        return f"{sp.min():.{prec}f}–{sp.max():.{prec}f}×"

    # The same for the SIFT1M re-timings (results/timing_controlled_sift1m.json).
    rts = {n: np.array([r["avg_query_time_ms"][n] for r in RETIME_S]) for n in RETIME_S[0]["avg_query_time_ms"]}
    rts_sp = {n: rts["linear_scan"] / t for n, t in rts.items()}
    rts_spread = {n: 100 * np.ptp(t) / np.median(t) for n, t in rts.items()}
    n_rts = len(rts["linear_scan"])
    rts_order = all(h_ < i_ < a_ for h_, i_, a_ in zip(rts["hnsw_ef32"], rts["ivf_nprobe16"],
                                                     rts["annoy_search_k5000"]))
    claim("SIFT1M re-timings: HNSW, IVF, Annoy order holds every round", rts_order,
          f"HNSW {rts['hnsw_ef32'].round(3)}, IVF {rts['ivf_nprobe16'].round(3)}")
    claim("SIFT1M re-timings: recall matches main run",
          all(abs(r["recall_at_10"][n] - v["recall_at_10"]) < 1e-9 for r in RETIME_S for n, v in
              (("hnsw_ef32", hn[S]), ("ivf_nprobe16", iv[S]), ("linear_scan", lin[S]),
               ("annoy_search_k5000", one("Annoy", S, "5000")))), "")
    # Medians that headline the abstract and conclusion, per dataset.
    rt_med = {S: (np.median(rts_sp["hnsw_ef32"]), np.median(rts_sp["ivf_nprobe16"])),
              W: (np.median(rt_sp["hnsw_ef64"]), np.median(rt_sp["ivf_nprobe16"]))}

    RETIME_LSH = (f"{rt_range('lsh_low', 2)} linear scan's speed across {NUM[n_rt]} re-timings, "
                  f"Section 3.11")
    RETIME_ANNOY = (f"In {NUM[n_rt]} interleaved re-timings Annoy ran at {rt_range('annoy_search_k100000', 2)} the "
                    f"speed of linear scan (Section 3.11), so it costs about as much as an exact search while still returning "
                    f"an approximate answer.")
    LSH_VS_LINEAR = "on Wikipedia it is no faster than linear scan at any setting"
    ANNOY_VS_LINEAR = "its near-exact setting costs about as much as linear scan on Wikipedia"
    # SIFT1M timing stability across the five k-means rebuilds (IVF, nprobe 16).
    # Recall spread (points) across k-means seeds, over every configuration tested.
    km_spread_all = []
    for ds in DS:
        for kind, cfgs in KMEANS[ds].items():
            if kind == "settings":
                continue
            per_seed = [cfgs[str(s)] for s in KMEANS["seeds"]]
            subs = [k for k in per_seed[0] if k.startswith("nprobe")] or [None]
            for sub in subs:
                v = [(p[sub] if sub else p)["recall_at_10"] for p in per_seed]
                km_spread_all.append(100 * (max(v) - min(v)))
    abl_k = sorted(ABLATION["euclidean"], key=int)
    abl_diff = [100 * (ABLATION["euclidean"][k]["recall_at_10"] - ABLATION["angular"][k]["recall_at_10"]) for k in abl_k]
    wiki_gap = [100 * (ABLATION["euclidean"][k]["recall_at_10"] - one("Annoy", W, k)["recall_at_10"]) for k in abl_k]
    misses_s = round((1 - lin[S]["recall_at_10"]) * 10 * nq[S])

    # IVF+PQ frontier: the sequence of m values it switches between.
    def frontier_ms(ds):
        seq = [pq_m(r["param"]) for r in pc.pareto(pts("IVF+PQ", ds))]
        return [m_ for i, m_ in enumerate(seq) if i == 0 or seq[i - 1] != m_]

    # Ranking by time to recall >= 0.9.
    def ranking(ds):
        reach = sorted((m for m in METHODS if fast[(m, ds)] and m != "Linear scan"),
                       key=lambda m: fast[(m, ds)]["avg_query_time_ms"])
        miss = sorted((m for m in METHODS if not fast[(m, ds)]),
                      key=lambda m: -best_recall(m, ds)["recall_at_10"])
        return reach, miss

    rk = {ds: ranking(ds) for ds in DS}

    # PQ-ADC vs IVF+PQ at equal m: (dataset, m, best IVF+PQ row, PQ-ADC row), and
    # the fastest IVF+PQ time that matches PQ-ADC's recall.
    pq_vs_ivfpq, match_ms = [], {}
    for ds in DS:
        for m_ in sorted({pq_m(r["param"]) for r in pts("PQ-ADC", ds)}):
            a_ = next(r for r in pts("PQ-ADC", ds) if pq_m(r["param"]) == m_)
            grp = [r for r in pts("IVF+PQ", ds) if pq_m(r["param"]) == m_]
            pq_vs_ivfpq.append((ds, m_, max(grp, key=lambda r: r["recall_at_10"]), a_))
            ok = [r["avg_query_time_ms"] for r in grp if r["recall_at_10"] >= a_["recall_at_10"]]
            match_ms[(ds, m_)] = min(ok) if ok else None
    n_wins = sum(b["recall_at_10"] > a["recall_at_10"] for _, _, b, a in pq_vs_ivfpq)

    # HNSW vs IVF time to recall >= 0.9: within 5% from a single timing run is
    # treated as a tie rather than a ranking.
    tied = {ds: abs(iv[ds]["avg_query_time_ms"] - hn[ds]["avg_query_time_ms"])
            / min(iv[ds]["avg_query_time_ms"], hn[ds]["avg_query_time_ms"]) < 0.05 for ds in DS}

    def hnsw_ivf(ds):
        # Headline speed-ups are re-timing medians (see RT_NOTE), not main-run values.
        sp_h, sp_i = rt_med[ds]
        if tied[ds]:
            return (f"HNSW and IVF are effectively tied, roughly {round((sp_h + sp_i) / 2, -1):.0f}× faster than "
                    f"linear scan")
        return (f"HNSW reaches recall@10 ≥ 0.9 fastest, about {fx(sp_h)} faster than linear scan against about "
                f"{fx(sp_i)} for IVF")

    ivf_hnsw_s = (iv[S]["avg_query_time_ms"] / hn[S]["avg_query_time_ms"],
                  np.median(rts["ivf_nprobe16"] / rts["hnsw_ef32"]))
    claim("IVF within about 2x of HNSW on SIFT1M (main and re-timed), tied on Wikipedia",
          all(r <= 2.25 for r in ivf_hnsw_s) and tied[W],
          f"SIFT1M ratios {ivf_hnsw_s[0]:.2f} (main), {ivf_hnsw_s[1]:.2f} (re-timed median)")
    RT_NOTE = f"(speed-ups are medians of {NUM[n_rts]} interleaved re-timings per dataset)"

    fastest_ivfpq = min(r["avg_query_time_ms"] for r in pts("IVF+PQ", S) + pts("IVF+PQ", W))

    # Highest recall at which a method other than HNSW/IVF is still on the global
    # Pareto frontier; above it, only HNSW and IVF are Pareto-optimal.
    global_cut = {}
    for ds in DS:
        front = pc.pareto([r for r in ROWS if r["dataset"] == ds and r["method"] != "Linear scan"])
        global_cut[ds] = max(r["recall_at_10"] for r in front if r["method"] not in ("HNSW", "IVF"))

    def sweep_growth(method, ds):
        by_param = sorted(pts(method, ds), key=lambda r: int(r["param"]))
        return by_param[-1]["avg_query_time_ms"] / by_param[0]["avg_query_time_ms"]

    sdc_top_w = max(pts("PQ-SDC", W), key=lambda r: pq_m(r["param"]))

    R = Report()
    R.title("Approximate Nearest Neighbour Search: An Empirical Comparison of Seven Methods "
            "on SIFT1M and Wikipedia Sentence Embeddings",
            ["SC4020 Project 1: Technical Review Report",
             "Group XX: Jacob, Andrea, Emma and Yuxuan",
             "October 2026"])

    # ------------------------------------------------------------------ Abstract
    R.h1("Abstract")
    R.p(f"Nearest-neighbour search over high-dimensional vectors is used in similarity search, recommendation "
        f"and retrieval-augmented systems, but exact search costs time linear in the database size. We compare "
        f"an exact linear scan with six approximate nearest-neighbour (ANN) methods: random-projection LSH, "
        f"product quantisation (PQ, with asymmetric and symmetric distances), the inverted file index (IVF), "
        f"IVF with PQ codes (IVF+PQ), HNSW and Annoy. We run them on the SIFT1M benchmark ({fint(nb[S])} × {sh[S]['base'][1]}) "
        f"and on {fint(nb[W])} Wikipedia sentence embeddings ({sh[W]['base'][1]}-dimensional). All methods run "
        f"single-threaded on one machine and are scored on recall@10, query time, build time and index size. "
        f"On SIFT1M, {hnsw_ivf(S)}; on Wikipedia, {hnsw_ivf(W)} {RT_NOTE}. IVF needs a fraction of HNSW's build time. "
        f"PQ-based indexes are the smallest ({rng(pq_sizes[S] + ivfpq_sizes[S], fmb)} MB on SIFT1M versus "
        f"{fmb(lin[S]['index_size_mb'])} MB of raw vectors) but cap recall: IVF+PQ, which matches plain PQ's recall "
        f"far faster and is more accurate in {n_wins} of {len(pq_vs_ivfpq)} code-size settings, peaks at "
        f"{fr(ivfpq_best[S]['recall_at_10'])} (SIFT1M) and "
        f"{fr(ivfpq_best[W]['recall_at_10'])} (Wikipedia). Our pure-Python LSH never reaches 0.9 on SIFT1M and is "
        f"no faster than brute force at any setting on Wikipedia. A controlled Annoy ablation shows that the distance "
        f"metric accounts for at most {max(abl_diff):.1f} recall points of the gap between the two datasets. "
        f"Retraining the k-means-based indexes with {len(KMEANS['seeds'])} seeds changes recall by at most "
        f"{max(km_spread_all):.1f} points, and a per-query analysis shows that IVF fails on queries whose true "
        f"neighbours are spread over many cells.")

    # ------------------------------------------------------------------ Introduction
    R.h1("1  Introduction")
    R.p("Given a database X = {x₁, …, xₙ} ⊂ ℝᵈ and a query q, k-nearest-neighbour (k-NN) search returns the k "
        "points closest to q under a distance such as Euclidean (L2) distance or cosine distance. It is the "
        "core operation behind content-based image retrieval, recommendation and semantic search, and more "
        "recently behind vector databases that store learned embeddings for retrieval-augmented generation.")
    R.p(f"Exact search compares the query with every database vector, costing O(nd) per query. On SIFT1M this is "
        f"{fint(nb[S])} distance computations in {sh[S]['base'][1]} dimensions, which our NumPy linear scan performs "
        f"in {fms(lin_ms[S])} ms per query on a single core. Cost grows linearly with both n and d, so exact search "
        f"does not scale to billion-vector collections or to services that must answer thousands of queries per "
        f"second, and classical space-partitioning trees lose their advantage in high dimensions (Indyk & Motwani, "
        f"1998). Approximate "
        f"nearest-neighbour (ANN) methods accept a small loss in accuracy, measured as recall against the exact "
        f"answer, in exchange for large reductions in query time or memory. Different methods buy this trade-off "
        f"with different structures, such as hash tables, compressed codes, inverted lists, graphs and trees, so "
        f"their behaviour depends on the data and on the operating point. ANN search is used at scale in "
        f"production: for example, Annoy was developed at Spotify for music recommendations (Bernhardsson, 2013), and libraries "
        f"such as FAISS (Johnson et al., 2021) serve embedding retrieval in industry.")
    R.p("**Contribution.** We implement or wrap six ANN methods and an exact baseline behind one interface "
        "(`build_index`, `search`), evaluate them with one shared harness on a standard benchmark (SIFT1M) and "
        "an applied dataset (Wikipedia sentence embeddings), and report:")
    R.bullets([
        "recall/latency trade-off curves and index sizes for all methods under identical, single-threaded "
        "conditions on one machine;",
        "the fastest configuration of each method that reaches recall@10 ≥ 0.9, and a ranking of methods;",
        "a parameter-sensitivity study of the search-time knobs (Annoy search_k, HNSW ef_search, IVF nprobe, "
        "PQ m, LSH configuration);",
        "an improvement study of IVF+PQ over plain PQ, and a controlled ablation isolating the effect of the "
        "distance metric on Annoy;",
        "a study of sensitivity to k-means initialisation, a per-query analysis of success and failure cases, "
        "and a summary of each method's strengths and weaknesses.",
    ])
    R.p("Section 2 describes the methods, Section 3 the experimental setup, results and discussion, and Section 4 "
        "concludes with practical recommendations.")

    # ------------------------------------------------------------------ Methods
    R.h1("2  Methods")
    R.p(f"All methods expose `build_index(vectors)` and `search(q, k)`, and share the evaluation harness in "
        f"`eval/harness.py`. We group the six ANN methods into three families: hashing (LSH), quantisation "
        f"(PQ, IVF, IVF+PQ) and graph/tree methods (HNSW, Annoy). LSH, PQ and IVF are covered in the course; "
        f"HNSW and Annoy go beyond it. {T('methods')} lists the implementation and parameters of each method.")

    R.h2("2.1  Exact baseline: linear scan")
    R.p("Linear scan computes the distance from the query to every database vector and returns the k smallest. "
        "Our implementation precomputes the squared norms ‖x‖² at build time and evaluates "
        "‖q − x‖² = ‖q‖² + ‖x‖² − 2 q·x with one matrix–vector product per query, followed by a partial sort "
        "(`argpartition`) for the top k. It defines both the exact answer and the latency every ANN method must beat.")

    R.h2("2.2  Hashing: random-projection LSH")
    sift_ws = [lsh_s[n]["bucket_width"] for n in lsh_s]
    wiki_ls = ", ".join(str(v) for v in sorted({c["num_tables"] for c in lsh_w.values()}))
    wiki_ws = ", ".join(f"{v:g}" for v in sorted({c["bucket_width"] for c in lsh_w.values()}))
    R.p(f"Locality-sensitive hashing (Indyk & Motwani, 1998; Datar et al., 2004) uses hash functions under which nearby points collide with higher "
        f"probability than distant ones. We implement the p-stable (E2LSH) scheme of Datar et al. (2004): each hash is "
        f"h(x) = ⌊(a·x + b) / w⌋ with a drawn from a standard Gaussian, b uniform in [0, w) and bucket width w. "
        f"K hashes are concatenated into one table key, and L independent tables are built. A query collects the "
        f"union of the buckets it falls into across the L tables (no multi-probing), then re-ranks these "
        f"candidates by exact L2 distance. Larger K makes buckets more selective, while larger L and w raise recall "
        f"at the cost of more candidates. Tables are Python dictionaries keyed by tuples of hash values, built in "
        f"NumPy chunks of {fint(lsh_mod.BUILD_CHUNK_SIZE)} vectors. On SIFT1M we use L = "
        f"{lsh_s['low']['num_tables']} and K = {lsh_s['low']['num_hashes']} with w ∈ "
        f"{{{', '.join(f'{w:g}' for w in sift_ws)}}}, i.e. "
        f"{', '.join(f'{w / lsh_mod.SIFT_PROJECTION_SCALE:.1f}' for w in sift_ws)} times the measured projection "
        f"standard deviation ({lsh_mod.SIFT_PROJECTION_SCALE:g}). On the unit-norm Wikipedia vectors we use "
        f"K = {lsh_w_low['num_hashes']} with L ∈ {{{wiki_ls}}} and w ∈ {{{wiki_ws}}}.")

    R.h2("2.3  Quantisation: PQ, IVF and IVF+PQ")
    R.p(f"**Product quantisation (PQ)** (Jégou et al., 2011) splits each d-dimensional vector into m sub-vectors of d/m dimensions "
        f"and learns a codebook of 2^nbits = {2 ** pqivf_mod.NBITS} centroids per sub-space with k-means. Each "
        f"database vector is stored as m one-byte centroid indices, so memory falls from 4d bytes to m bytes. "
        f"Distances are estimated from lookup tables. With **asymmetric distance computation (ADC)** the query "
        f"stays unquantised: per query, an m × {2 ** pqivf_mod.NBITS} table of query-to-centroid distances is built, "
        f"and each database distance is a sum of m lookups. With **symmetric distance computation (SDC)** the "
        f"query is also quantised, and distances come from precomputed centroid-to-centroid tables. SDC adds "
        f"quantisation error on the query side, which is why ADC is normally preferred (Jégou et al., 2011). We use FAISS "
        f"`IndexPQ` (Johnson et al., 2021), which scans all codes exhaustively, with m ∈ {{{', '.join(map(str, pqivf_mod.SIFT_CFG['pq_ms']))}}} "
        f"on SIFT1M and m ∈ {{{', '.join(map(str, pqivf_mod.WIKI_CFG['pq_ms']))}}} on Wikipedia (m must divide d).")
    R.p(f"**Inverted file index (IVF)** partitions the space with k-means into nlist cells (nlist = "
        f"{pqivf_mod.SIFT_CFG['nlist']} for SIFT1M and {pqivf_mod.WIKI_CFG['nlist']} for Wikipedia). Each vector "
        f"is stored, uncompressed, in the list of its nearest centroid. A query visits only the nprobe cells whose "
        f"centroids are nearest and computes exact distances within them, so nprobe trades recall for speed "
        f"(swept over {{{', '.join(map(str, pqivf_mod.NPROBES))}}}). We use FAISS `IndexIVFFlat`.")
    R.p("**IVF+PQ** (Jégou et al., 2011) combines the two: IVF selects the cells, and each vector is stored as a PQ code of its "
        "residual x − c(x) with respect to its cell centroid c(x). Residuals have smaller spread than raw vectors, "
        "so the same m bytes encode them more accurately. We use FAISS `IndexIVFPQ` with the same nlist and m "
        "values as above and no exact re-ranking of candidates. All quantisers (k-means centroids and PQ "
        "codebooks) are trained on the separate learn/train set, never on the base vectors.")

    R.h2("2.4  Graphs and trees: HNSW and Annoy")
    R.p(f"**HNSW** (Malkov & Yashunin, 2020) builds a hierarchy of proximity graphs. Each point is inserted at a random maximum layer "
        f"with exponentially decaying probability and linked to up to M neighbours per layer (2M on the bottom "
        f"layer in hnswlib; candidates found with a "
        f"beam of width efConstruction). Search descends greedily from the sparse top layer and runs a beam search "
        f"of width ef_search on the bottom layer, so ef_search is the recall/speed knob. We use hnswlib with "
        f"M = {hnsw_mod.M} and efConstruction = {hnsw_mod.EF_CONSTRUCTION}, sweeping ef_search over "
        f"{{{', '.join(map(str, hnsw_mod.EF_SEARCH_SWEEP))}}}. Distances are L2 on SIFT1M and inner product on the "
        f"normalised Wikipedia vectors.")
    R.p(f"**Annoy** (Bernhardsson, 2013) builds a forest of n_trees random-projection trees. Each internal node splits its points "
        f"by the hyperplane equidistant from two sampled points, until leaves are small. At query time all trees "
        f"are searched together with a priority queue until search_k nodes have been inspected; the collected "
        f"candidates are then ranked by exact distance. We use n_trees = {annoy_mod.N_TREES} and sweep "
        f"search_k over {{{', '.join(fint(k) for k in annoy_mod.SEARCH_K_SWEEP)}}}, with the euclidean metric on "
        f"SIFT1M and the angular metric on Wikipedia. The index is saved to disk and memory-mapped, which is what "
        f"makes Annoy easy to share between processes.")

    R.table("methods", "Methods, implementations and parameters. The last column is the search-time parameter swept "
            "to trace each recall/latency curve.",
            ["Method", "Family", "Implementation", "Fixed parameters", "Swept parameter"],
            [["Linear scan", "Exact", "NumPy", "none", "none"],
             ["LSH", "Hashing", "NumPy + Python dict", "L tables, K hashes, width w", "3–4 (L, K, w) configs"],
             ["PQ (ADC, SDC)", "Quantisation", "FAISS IndexPQ", f"nbits = {pqivf_mod.NBITS}", "m (bytes per vector)"],
             ["IVF", "Quantisation", "FAISS IndexIVFFlat",
              f"nlist = {pqivf_mod.SIFT_CFG['nlist']} / {pqivf_mod.WIKI_CFG['nlist']}", "nprobe"],
             ["IVF+PQ", "Quantisation", "FAISS IndexIVFPQ",
              f"nlist as IVF, nbits = {pqivf_mod.NBITS}", "m × nprobe"],
             ["HNSW", "Graph", "hnswlib", f"M = {hnsw_mod.M}, efC = {hnsw_mod.EF_CONSTRUCTION}", "ef_search"],
             ["Annoy", "Tree", "annoy", f"n_trees = {annoy_mod.N_TREES}", "search_k"]],
            align=["l", "l", "l", "l", "l"],
            note="nlist values are given as SIFT1M / Wikipedia.")

    # ------------------------------------------------------------------ Experiments
    R.h1("3  Experiments")
    R.h2("3.1  Datasets")
    R.p(f"{T('data')} summarises the two datasets. **SIFT1M** (the ANN_SIFT1M set of the TEXMEX corpus; Jégou et al., 2011) "
        f"contains SIFT descriptors (Lowe, 2004) with a provided exact ground truth of the {sh[S]['gt'][1]} nearest base "
        f"vectors per query, of which we use the first 10. Distances are L2. The **Wikipedia** dataset consists of "
        f"English Wikipedia sentences drawn from the Hugging Face dataset `sentence-transformers/wikipedia-en-sentences` "
        f"(7.87 million pre-split sentences) (Sentence Transformers, n.d.), embedded with the sentence-transformers "
        f"model `{wcfg['model']}` (Reimers & Gurevych, 2019). The model name and split sizes are recorded in the `config` field of our "
        f"`data/wiki/wiki_sentences.json`. The base, query and train sets were sampled with seed {wcfg['seed']} "
        f"and are disjoint: "
        + ("no sentence appears in more than one of them, so no query is also a base vector. "
           if wcfg["overlap"] == 0 else f"[CHECK: {wcfg['overlap']} sentences are shared between splits]. ")
        + f"The sampling is implemented in `scripts/generate_embeddings.py`, which shuffles the corpus with seed "
          f"{wcfg['seed']} and takes three non-overlapping slices. "
        + f"All Wikipedia vectors are unit-norm: the embedding script normalises them, and each method normalises "
        f"them again (Annoy instead uses its angular metric, which is norm-independent), so L2 distance ranks "
        f"neighbours identically to cosine distance "
        f"(‖q − x‖² = 2 − 2 cos(q, x) for unit vectors). Exact ground truth is computed by brute force on the "
        f"normalised vectors (by the shared harness, and with an equivalent FAISS flat index in the PQ/IVF "
        f"script). On both datasets the learn/train split is used only to train PQ and IVF quantisers.")
    R.table("data", "Datasets. Counts are read from the data files.",
            ["Dataset", "Base", "Queries", "Learn / train", "Dim.", "Metric", "Ground truth"],
            [["SIFT1M", fint(nb[S]), fint(nq[S]), fint(sh[S]["train"][0]), sh[S]["base"][1], "L2",
              f"provided (top {sh[S]['gt'][1]})"],
             ["Wikipedia", fint(nb[W]), fint(nq[W]), fint(sh[W]["train"][0]), sh[W]["base"][1],
              "cosine (L2 on unit vectors)", "brute force"]],
            align=["l", "r", "r", "r", "r", "l", "l"], widths=[1.9, 1.9, 1.6, 2.0, 1.1, 3.5, 3.6])

    R.h2("3.2  Protocol and setup")
    R.p("We report four measures for every configuration:")
    R.bullets([
        "**Recall@10**: the fraction of the 10 true nearest neighbours that appear in the 10 returned ids, "
        "averaged over all queries.",
        "**Query time**: mean wall-clock time of one `search` call over the whole query set, with queries issued "
        "one at a time (no batching) and timed individually with `time.perf_counter`, Python call overhead "
        "included.",
        "**Build time**: wall-clock time of `build_index`, including quantiser training for PQ and IVF.",
        "**Index size**: size of the serialised index on disk. For linear scan and LSH this includes the stored raw "
        "vectors, and the LSH figure also includes its pickled hash tables.",
    ])
    env = environment()
    R.p(f"All runs used one machine: an {env['cpu']} with {env['cores']} cores and {env['mem_gb']} GB of memory, "
        f"running macOS {env['macos']}, Python {env['python']}, NumPy {env['numpy']}, faiss-cpu {env['faiss-cpu']}, "
        f"hnswlib {env['hnswlib']} and annoy {env['annoy']}. To make query times comparable, "
        "every method ran single-threaded: `faiss.omp_set_num_threads(1)` in the PQ/IVF script, "
        "`set_num_threads(1)` on the hnswlib index, and `n_jobs=1` for the Annoy build, with "
        "OMP_NUM_THREADS = VECLIB_MAXIMUM_THREADS = OPENBLAS_NUM_THREADS = MKL_NUM_THREADS = 1 set for every run "
        "so that NumPy's BLAS also used one core (the full run commands are recorded in the repository README). Annoy's `save` re-opens the index as a memory-mapped file, so the first "
        "searches after a build page the index in from disk. We therefore run one untimed warm-up pass over all "
        "queries at the smallest search_k before timing; without it, the first timed search_k on Wikipedia was "
        "inflated by disk reads. The other "
        "indexes live in memory and need no warm-up. Each configuration was timed once over the full query set, so the "
        "main results come from single runs; Section 3.11 reports repeated, interleaved timings of the headline "
        "configurations.")

    R.h2("3.3  Overall results")
    order = {ds: join_and([f"{m} ({fms(fast[(m, ds)]['avg_query_time_ms'])} ms)" for m in rk[ds][0]]) for ds in DS}
    R.p(f"{F('sift')} and {F('wiki')} plot recall@10 against query time for every configuration, and "
        f"{T('summary')} lists, for each method, its best recall and its fastest configuration reaching "
        f"recall@10 ≥ 0.9. On SIFT1M the methods that reach 0.9 rank {order[S]}, against {fms(lin_ms[S])} ms for "
        f"linear scan; {join_and(rk[S][1])} never reach it. On Wikipedia the order is {order[W]}, against "
        f"{fms(lin_ms[W])} ms for linear scan, so LSH reaches 0.9 only at "
        f"{fx(lshw['avg_query_time_ms'] / lin_ms[W])} the cost of brute force; {join_and(rk[W][1])} never reach 0.9."
        + (f" HNSW and IVF differ there by under 5% in the main run and swap order across re-timings (IVF faster "
           f"in {NUM[int((rt['ivf_nprobe16'] < rt['hnsw_ef64']).sum())]} of {NUM[n_rt]}), so we treat them as tied "
           f"rather than ranked." if tied[W] else ""))
    R.p(f"HNSW and IVF form the best trade-off on both datasets whenever recall above "
        f"{fr(max(global_cut[S], global_cut[W]))} is required. Below that, only IVF+PQ is faster, at the cost of "
        f"lower recall. HNSW reaches 0.9 in {fms(hn[S]['avg_query_time_ms'])} ms on SIFT1M "
        f"({fx(lin_ms[S] / hn[S]['avg_query_time_ms'])} in the main run; median {fx(rt_med[S][0])} across "
        f"{NUM[n_rts]} re-timings, in each of which it was the fastest) and {fms(hn[W]['avg_query_time_ms'])} "
        f"ms on Wikipedia ({fx(lin_ms[W] / hn[W]['avg_query_time_ms'])} in the main run; median "
        f"{fx(np.median(rt_sp['hnsw_ef64']))} across {NUM[n_rt]} re-timings). IVF takes about "
        f"{fx(iv[S]['avg_query_time_ms'] / hn[S]['avg_query_time_ms'])} HNSW's time on SIFT1M (median "
        f"{fx(np.median(rts['ivf_nprobe16'] / rts['hnsw_ef32']))} in the re-timings) and is nearly tied on "
        f"Wikipedia ({fms(iv[W]['avg_query_time_ms'])} ms against {fms(hn[W]['avg_query_time_ms'])} ms), "
        f"while building {fx(hnsw_build[S] / ivf_build[S])} and {fx(hnsw_build[W] / ivf_build[W])} faster "
        f"({fs(ivf_build[S])} s against {fs(hnsw_build[S])} s on SIFT1M). At the high-recall end both exceed 0.99 "
        f"(HNSW {fr(best_recall('HNSW', S)['recall_at_10'])} / {fr(best_recall('HNSW', W)['recall_at_10'])}, IVF "
        f"{fr(best_recall('IVF', S)['recall_at_10'])} / {fr(best_recall('IVF', W)['recall_at_10'])} on SIFT1M / "
        f"Wikipedia) in {rng([best_recall(m, ds)['avg_query_time_ms'] for m in ('HNSW', 'IVF') for ds in DS], fms)} ms. "
        f"Annoy follows: it reaches 0.9 at {fms(an[S]['avg_query_time_ms'])} ms on SIFT1M and "
        f"{fms(an[W]['avg_query_time_ms'])} ms on Wikipedia.")

    # Methods with no configuration on the overall (all-method) Pareto frontier.
    def dominated(ds, methods):
        front = {(r["method"], r["param"]) for r in pc.pareto([r for r in ROWS if r["dataset"] == ds])}
        return all(not any(m == fm for fm, _ in front) for m in methods)

    for key, ds in (("sift", S), ("wiki", W)):
        R.figure(key, os.path.join(FIG_DIR, f"recall_vs_latency_{ds}.png"), width_cm=14.8, caption=
                 f"Recall@10 against mean query time on {DS_NAME[ds]} (log scale; up and to the left is better). "
                 f"Each line joins a method's Pareto-optimal configurations; faint markers are dominated "
                 f"configurations. The star and dashed line mark exact linear scan ({fms(lin_ms[ds])} ms): any "
                 f"point to the right of the line is slower than brute force. "
                 + ("Note the flat IVF+PQ curve (compression ceiling)"
                    + (", and that every LSH and PQ configuration is beaten by another method on both speed and "
                       "recall." if dominated(S, ("LSH", "PQ-ADC", "PQ-SDC")) else ".") if ds == S else
                    "Note that no LSH configuration is meaningfully faster than linear scan, and that HNSW and IVF "
                    "nearly coincide."))

    sum_rows = []
    for ds in DS:
        sum_rows.append(DS_NAME[ds])
        for m in METHODS:
            p = pts(m, ds)
            f_ = fast[(m, ds)]
            sum_rows.append([
                m, fr(best_recall(m, ds)["recall_at_10"]),
                pretty(m, ds, f_["param"]) if f_ else "not reached",
                fms(f_["avg_query_time_ms"]) if f_ else "n/a",
                fx(lin_ms[ds] / f_["avg_query_time_ms"]) if f_ else "n/a",
                rng([r["build_time_sec"] for r in p], fs),
                rng([r["index_size_mb"] for r in p], fmb)])
    R.table("summary", f"Summary per method. Speed-up is relative to linear scan on the same dataset, from the main "
            f"run. Re-timings (Section 3.11) give lower medians for HNSW and IVF: {fx(rt_med[S][0])} and "
            f"{fx(rt_med[S][1])} on SIFT1M, {fx(rt_med[W][0])} and {fx(rt_med[W][1])} on Wikipedia. Build time and "
            "index size are ranges where a method has several index configurations.",
            ["Method", "Max recall@10", "Fastest config with recall@10 ≥ 0.9", "Time (ms)", "Speed-up",
             "Build (s)", "Index (MB)"], sum_rows, align=["l", "r", "l", "r", "r", "r", "r"],
            widths=[2.1, 1.9, 4.4, 1.6, 1.6, 1.7, 1.9])

    R.h2("3.4  Memory")
    R.p(f"{F('size')} compares index sizes. Methods that keep the raw float32 vectors cannot be smaller than the data "
        f"({fmb(lin[S]['index_size_mb'])} MB for SIFT1M, {fmb(lin[W]['index_size_mb'])} MB for Wikipedia): IVF adds "
        f"only ids and centroids ({fmb(ivf_size[S])} / {fmb(ivf_size[W])} MB), HNSW adds its graph "
        f"({fmb(pts('HNSW', S)[0]['index_size_mb'])} / {fmb(pts('HNSW', W)[0]['index_size_mb'])} MB), and Annoy is "
        f"the largest at {fmb(pts('Annoy', S)[0]['index_size_mb'])} / {fmb(pts('Annoy', W)[0]['index_size_mb'])} MB, "
        f"because its {annoy_mod.N_TREES} trees store split hyperplanes in addition to the items. PQ and IVF+PQ store "
        f"only codes: {rng(pq_sizes[S], fmb)} MB and {rng(ivfpq_sizes[S], fmb)} MB on SIFT1M, i.e. "
        f"{fx(lin[S]['index_size_mb'] / max(ivfpq_sizes[S]))}–{fx(lin[S]['index_size_mb'] / min(pq_sizes[S]))} "
        f"smaller than the raw vectors. IVF+PQ is larger than PQ at equal m (about "
        f"{round(np.mean([min(ivfpq_sizes[d]) / min(pq_sizes[d]) for d in DS])):.0f}× at m = "
        f"{min(pqivf_mod.SIFT_CFG['pq_ms'])}; see Section 3.7) because it also stores an 8-byte id per vector in its "
        f"inverted lists.")
    R.figure("size", os.path.join(FIG_DIR, "index_size.png"),
             "Serialised index size per method (log scale). Bars show the smallest configuration and whiskers the "
             "largest, where configurations differ (PQ and IVF+PQ by m, LSH by number of tables). Note the "
             "gap of one to two orders of magnitude between the compressed PQ-family indexes and everything else.")

    R.h2("3.5  Parameter sensitivity")
    sweep_rows = []
    for method, label, key, params in (
            ("HNSW", "ef_search", int, hnsw_mod.EF_SEARCH_SWEEP),
            ("Annoy", "search_k", int, annoy_mod.SEARCH_K_SWEEP),
            ("IVF", "nprobe", None, pqivf_mod.NPROBES)):
        sweep_rows.append(f"{method} ({label})")
        for v in params:
            cells = [fint(v) if method == "Annoy" else str(v)]
            for ds in DS:
                r = (one(method, ds, str(v)) if method != "IVF" else
                     next(x for x in pts("IVF", ds) if ivf_nprobe(x["param"]) == v))
                cells += [fr(r["recall_at_10"]), fms(r["avg_query_time_ms"])]
            sweep_rows.append(cells)
    R.p(f"{T('sweeps')} lists the search-time sweeps of HNSW, Annoy and IVF. All three trade recall for time "
        f"smoothly and monotonically. Recall rises steeply at first and then saturates, while query time keeps "
        f"growing (by {fx(sweep_growth('HNSW', S))} for the {fx(max(hnsw_mod.EF_SEARCH_SWEEP) / min(hnsw_mod.EF_SEARCH_SWEEP))} "
        f"range of ef_search on SIFT1M), so each doubling of the parameter buys less recall. For IVF, nprobe = "
        f"{ivf_nprobe(iv[S]['param'])} (searching {100 * ivf_nprobe(iv[S]['param']) / pqivf_mod.SIFT_CFG['nlist']:.1f}% "
        f"of the {pqivf_mod.SIFT_CFG['nlist']} cells) already gives {fr(iv[S]['recall_at_10'])} on SIFT1M. "
        f"HNSW at ef_search = {hnsw_mod.EF_SEARCH_SWEEP[0]} starts at {fr(one('HNSW', S, str(hnsw_mod.EF_SEARCH_SWEEP[0]))['recall_at_10'])} "
        f"(SIFT1M) and {fr(one('HNSW', W, str(hnsw_mod.EF_SEARCH_SWEEP[0]))['recall_at_10'])} (Wikipedia).")
    ad = ANNOY_DEFAULT
    R.p(f"Search-time parameters matter as much as the choice of method. Annoy illustrates this: when search_k is "
        f"not given, the library inspects only n_trees × k = {ad[S]['search_k_effective']} nodes. Measured at this "
        f"default, Annoy reaches only {fr(ad[S]['recall_at_10'])} on SIFT1M and {fr(ad[W]['recall_at_10'])} on "
        f"Wikipedia ({fms(ad[S]['avg_query_time_ms'])} ms and {fms(ad[W]['avg_query_time_ms'])} ms per query), "
        f"against {fr(annoy_max[S]['recall_at_10'])} and {fr(annoy_max[W]['recall_at_10'])} at "
        f"search_k = {fint(max(annoy_mod.SEARCH_K_SWEEP))}; doubling the budget to "
        f"{fint(min(annoy_mod.SEARCH_K_SWEEP))} already lifts recall to {fr(annoy_min[S]['recall_at_10'])} and "
        f"{fr(annoy_min[W]['recall_at_10'])}. An Annoy deployment left at its default would therefore look far worse "
        f"than the method is. (The default was measured on a separate build of the same configuration, "
        f"`results/annoy_default_search_k.json`; Annoy's trees are random, so recall differs slightly between "
        f"builds.) The same holds for HNSW's ef_search and IVF's nprobe: the index is built once, and the "
        f"query-time parameter alone moves it along the whole curve.")
    R.table("sweeps", "Search-time parameter sweeps (one index per method and dataset). Recall@10 and mean query "
            "time in ms.",
            ["Parameter", "SIFT1M recall", "SIFT1M ms", "Wikipedia recall", "Wikipedia ms"], sweep_rows,
            split_ok=True)

    lsh_rows = []
    for ds in DS:
        lsh_rows.append(DS_NAME[ds])
        for n, c in LSH_RAW[ds].items():
            lsh_rows.append([n, c["num_tables"], c["num_hashes"], f"{c['bucket_width']:g}",
                             f"{c['candidate_pct']:.1f}%", fr(c["recall_at_10"]), fms(c["avg_query_time_ms"]),
                             fs(c["build_time_sec"]), fmb(c["index_size_mb"])])
    R.p(f"{T('lsh')} shows the LSH configurations. Recall is governed by how many candidates the buckets return. "
        f"On SIFT1M, widening w from {lsh_s['low']['bucket_width']:g} to {lsh_s['high']['bucket_width']:g} raises "
        f"the candidate set from {lsh_s['low']['candidate_pct']:.1f}% to {lsh_s['high']['candidate_pct']:.1f}% of "
        f"the base and recall from {fr(lsh_s['low']['recall_at_10'])} to {fr(lsh_s['high']['recall_at_10'])}, and "
        f"query time grows with the candidate count. On Wikipedia, reaching {fr(lsh_w['very_high']['recall_at_10'])} "
        f"requires inspecting {lsh_w['very_high']['candidate_pct']:.1f}% of the base. Our earlier Wikipedia "
        f"configuration used K = {OLD_LSH_W[0]['num_hashes']} hashes per table (commit `{OLD_LSH_COMMIT}`); its "
        f"buckets were so wide that its three settings re-ranked "
        f"{join_and([f'{c['candidate_pct']:.1f}%' for c in OLD_LSH_W])} "
        f"of the base (recall {join_and([fr(c['recall_at_10']) for c in OLD_LSH_W])}), approaching a full scan at the "
        f"high end, which motivated the switch to K = {lsh_w_low['num_hashes']}. We quote only its candidate "
        f"fractions and recall, as that run's timings were measured on a different machine.")
    R.table("lsh", "LSH configurations: L tables, K hashes per table, bucket width w. Candidates is the mean share of "
            "the base vectors re-ranked per query.",
            ["Config", "L", "K", "w", "Candidates", "Recall@10", "ms", "Build (s)", "Index (MB)"], lsh_rows)

    pq_rows = []
    for ds in DS:
        pq_rows.append(DS_NAME[ds])
        for m_ in sorted({pq_m(r["param"]) for r in pts("PQ-ADC", ds)}):
            a_ = next(r for r in pts("PQ-ADC", ds) if pq_m(r["param"]) == m_)
            s_ = next(r for r in pts("PQ-SDC", ds) if pq_m(r["param"]) == m_)
            pq_rows.append([m_, f"{sh[ds]['base'][1] // m_}", fr(a_["recall_at_10"]), fms(a_["avg_query_time_ms"]),
                            fr(s_["recall_at_10"]), fms(s_["avg_query_time_ms"]), fs(a_["build_time_sec"]),
                            fmb(a_["index_size_mb"])])
    pq_top = {ds: max(pts("PQ-ADC", ds), key=lambda r: pq_m(r["param"])) for ds in DS}
    R.p(f"For PQ ({T('pq')}), m sets both the code size (m bytes per vector) and the accuracy. On SIFT1M, ADC "
        f"recall rises from {fr(min(r['recall_at_10'] for r in pts('PQ-ADC', S)))} at m = "
        f"{min(pqivf_mod.SIFT_CFG['pq_ms'])} to {fr(pq_top[S]['recall_at_10'])} at m = {pq_m(pq_top[S]['param'])}. "
        f"On Wikipedia, m = {min(pqivf_mod.WIKI_CFG['pq_ms'])} (48 dimensions per sub-quantiser) gives only "
        f"{fr(min(r['recall_at_10'] for r in pts('PQ-ADC', W)))}, and m = {pq_m(pq_top[W]['param'])} is needed to "
        f"reach {fr(pq_top[W]['recall_at_10'])}. Because `IndexPQ` scans every code, query time also grows with m, "
        f"and on Wikipedia the largest m approaches the cost of linear scan "
        f"({fms(pq_top[W]['avg_query_time_ms'])} ms against {fms(lin_ms[W])} ms). SDC is less accurate and slower "
        f"than ADC at every m.")
    R.table("pq", "PQ by number of sub-quantisers m (nbits = 8, so m bytes per vector). Both distance modes use the "
            "same index.",
            ["m", "Dims / sub-q.", "ADC recall", "ADC ms", "SDC recall", "SDC ms", "Build (s)", "Index (MB)"], pq_rows)

    R.h2("3.6  Sensitivity to k-means initialisation")
    km_seeds = KMEANS["seeds"]
    km_rows, km_spread = [], {}

    def km_vals(ds, kind, sub=None):
        return np.array([(KMEANS[ds][kind][str(s)][sub] if sub else KMEANS[ds][kind][str(s)])["recall_at_10"]
                         for s in km_seeds])

    for ds in DS:
        st = KMEANS[ds]["settings"]
        km_rows.append(DS_NAME[ds])
        cfgs = [(f"IVF, nprobe {n}", "ivf", f"nprobe{n}") for n in st["ivf_nprobe"]]
        cfgs += [(f"PQ-ADC, m {st['pq_m']}", "pq", None),
                 (f"IVF+PQ, m {st['ivfpq_m']}, nprobe {st['ivfpq_nprobe']}", "ivfpq", None)]
        for label, kind, sub in cfgs:
            v = km_vals(ds, kind, sub)
            km_spread[(ds, label)] = 100 * (v.max() - v.min())
            km_rows.append([label, fr(v.min()), fr(v.max()), f"{100 * (v.max() - v.min()):.2f}",
                            f"{100 * v.std(ddof=1):.2f}"])
    # The default seed must reproduce the main IVF run exactly.
    km_repro = all(abs(KMEANS[ds]["ivf"]["1234"]["nprobe16"]["recall_at_10"]
                       - next(r for r in pts("IVF", ds) if ivf_nprobe(r["param"]) == 16)["recall_at_10"]) < 1e-9
                   for ds in DS)
    wiki16 = km_vals(W, "ivf", "nprobe16")
    imb = {ds: [KMEANS[ds]["ivf"][str(s)]["imbalance_factor"] for s in km_seeds] for ds in DS}
    pq_sample = 256 * 2 ** pqivf_mod.NBITS
    R.p(f"IVF learns its coarse centroids, and PQ its codebooks, with k-means, whose result depends on the "
        f"random initial centroids. FAISS seeds k-means with a fixed default (1234), so every IVF, PQ and IVF+PQ "
        f"result above comes from a single initialisation. To measure how much this matters, we retrained the "
        f"three indexes with seeds {{{', '.join(map(str, km_seeds))}}} at fixed operating points, changing nothing "
        f"else ({T('kmeans')}; `results/kmeans_seed_sensitivity.py`). FAISS caps k-means at 256 training points "
        f"per centroid, so the IVF quantisers use all {fint(sh[S]['train'][0])} training vectors, while each PQ "
        f"codebook (256 centroids) trains on a seed-dependent sample of {fint(pq_sample)}."
        + (" The default seed reproduces the main results exactly." if km_repro else
           " [CHECK: the default seed did not reproduce the main results].")
        + f" Initialisation has little effect: across seeds, recall varies by at most "
        f"{max(v for (ds, _), v in km_spread.items() if ds == S):.2f} points on SIFT1M and "
        f"{max(v for (ds, _), v in km_spread.items() if ds == W):.2f} on Wikipedia, far less than the gaps between "
        f"methods. It is largest for IVF with a single probe, where recall depends most on exactly where cell "
        f"boundaries fall, and smallest for IVF+PQ. The cells stay balanced for every seed (imbalance factor "
        f"{min(imb[S]):.3f}–{max(imb[S]):.3f} on SIFT1M and {min(imb[W]):.3f}–{max(imb[W]):.3f} on Wikipedia, "
        f"where 1 is perfectly even). One result is seed-sensitive, however: on Wikipedia, IVF at nprobe = 16 "
        f"ranges from {fr(wiki16.min())} to {fr(wiki16.max())}, so {int((wiki16 < 0.9).sum())} of "
        f"{len(km_seeds)} seeds fall below 0.9. Whether IVF reaches the 0.9 threshold at nprobe = 16 there is "
        f"therefore partly a matter of initialisation, which is further reason to treat HNSW and IVF on Wikipedia "
        f"as tied.")
    R.table("kmeans", f"Recall@10 across {len(km_seeds)} k-means seeds (one index per seed), with spread and "
            "standard deviation in recall points (×100).",
            ["Configuration", "Min recall", "Max recall", "Range (pts)", "Std. dev. (pts)"], km_rows)

    R.h2("3.7  Improvement study: PQ versus IVF+PQ")
    ivfpq_rows = []
    for ds in DS:
        ivfpq_rows.append(DS_NAME[ds])
        for m_ in sorted({pq_m(r["param"]) for r in pts("PQ-ADC", ds)}):
            a_ = next(r for r in pts("PQ-ADC", ds) if pq_m(r["param"]) == m_)
            grp = [r for r in pts("IVF+PQ", ds) if pq_m(r["param"]) == m_]
            best = max(grp, key=lambda r: r["recall_at_10"])
            match = [r for r in grp if r["recall_at_10"] >= a_["recall_at_10"]]
            mt = min(match, key=lambda r: r["avg_query_time_ms"]) if match else None
            ivfpq_rows.append([m_, fr(a_["recall_at_10"]), fms(a_["avg_query_time_ms"]),
                               f"{fr(best['recall_at_10'])} ({ivf_nprobe(best['param'])})",
                               f"{fms(mt['avg_query_time_ms'])} ({ivf_nprobe(mt['param'])})" if mt else "n/a",
                               fx(a_["avg_query_time_ms"] / mt["avg_query_time_ms"]) if mt else "n/a",
                               fmb(a_["index_size_mb"]), fmb(grp[0]["index_size_mb"])])
    wins = [(ds, m_, b, a) for ds, m_, b, a in pq_vs_ivfpq if b["recall_at_10"] > a["recall_at_10"]]
    losses = [(ds, m_, b, a) for ds, m_, b, a in pq_vs_ivfpq if b["recall_at_10"] <= a["recall_at_10"]]
    loss_txt = join_and([f"{DS_NAME[ds]} at m = {m_} (IVF+PQ peaks at {fr(b['recall_at_10'])} against PQ's "
                         f"{fr(a['recall_at_10'])})" for ds, m_, b, a in losses])
    for ds, cfg in ((S, pqivf_mod.SIFT_CFG), (W, pqivf_mod.WIKI_CFG)):
        extra = (nb[ds] * 8 + cfg["nlist"] * sh[ds]["base"][1] * 4) / 2 ** 20  # ids + float32 centroids
        claim(f"IVF+PQ size overhead = 8-byte ids + centroids ({ds})",
              all(abs(b_ - a_ - extra) < 0.02 * extra for a_, b_ in zip(sorted(pq_sizes[ds]), sorted(set(ivfpq_sizes[ds])))),
              f"expected {extra:.2f} MB; measured "
              + ", ".join(f"{b_ - a_:.2f}" for a_, b_ in zip(sorted(pq_sizes[ds]), sorted(set(ivfpq_sizes[ds])))))
    R.p(f"IVF+PQ adds a coarse IVF partition to PQ. {T('ivfpq')} compares the two at equal code size m. IVF+PQ "
        f"reaches higher recall in {len(wins)} of the {len(pq_vs_ivfpq)} settings, as expected from encoding "
        f"residuals rather than raw vectors"
        + (f"; the exception is {loss_txt}" if losses else "")
        + f". Wherever IVF+PQ matches PQ's recall it does so much faster, because it scans only nprobe lists "
        f"instead of all codes: at the largest m, plain PQ needs {fms(pq_top[S]['avg_query_time_ms'])} ms (SIFT1M) and "
        f"{fms(pq_top[W]['avg_query_time_ms'])} ms (Wikipedia) for recall that IVF+PQ matches in "
        f"{fms(match_ms[(S, pq_m(pq_top[S]['param']))])} ms and {fms(match_ms[(W, pq_m(pq_top[W]['param']))])} ms. "
        f"The cost is index size: IVF+PQ stores an 8-byte id per vector next to its m-byte code, plus the coarse "
        f"centroids. That roughly doubles the index at m = {min(pqivf_mod.SIFT_CFG['pq_ms'])} "
        f"({fmb(min(pq_sizes[S]))} against {fmb(min(ivfpq_sizes[S]))} MB on SIFT1M, "
        f"{fmb(min(pq_sizes[W]))} against {fmb(min(ivfpq_sizes[W]))} MB on Wikipedia) but matters less as codes "
        f"grow: at the largest m it adds {100 * (max(ivfpq_sizes[S]) / max(pq_sizes[S]) - 1):.0f}% on SIFT1M "
        f"({fmb(max(pq_sizes[S]))} against {fmb(max(ivfpq_sizes[S]))} MB) and "
        f"{100 * (max(ivfpq_sizes[W]) / max(pq_sizes[W]) - 1):.0f}% on Wikipedia ({fmb(max(pq_sizes[W]))} against "
        f"{fmb(max(ivfpq_sizes[W]))} MB).")
    def at_np(ds, nprobe):
        m_top = pq_m(ivfpq_best[ds]["param"])
        return next(r for r in pts("IVF+PQ", ds) if pq_m(r["param"]) == m_top and ivf_nprobe(r["param"]) == nprobe)
    np_top = max(pqivf_mod.NPROBES)
    R.p(f"The improvement has a ceiling. For the largest m, raising nprobe from 32 to {np_top} adds only "
        f"{pts_diff(ivfpq_best[S]['recall_at_10'], at_np(S, 32)['recall_at_10'])} recall points on SIFT1M and "
        f"{pts_diff(ivfpq_best[W]['recall_at_10'], at_np(W, 32)['recall_at_10'])} on Wikipedia while multiplying "
        f"query time by {fx(ivfpq_best[S]['avg_query_time_ms'] / at_np(S, 32)['avg_query_time_ms'])} and "
        f"{fx(ivfpq_best[W]['avg_query_time_ms'] / at_np(W, 32)['avg_query_time_ms'])}. IVF+PQ "
        f"saturates at {fr(ivfpq_best[S]['recall_at_10'])} on SIFT1M and {fr(ivfpq_best[W]['recall_at_10'])} on "
        f"Wikipedia, whereas IVF with uncompressed vectors reaches {fr(best_recall('IVF', S)['recall_at_10'])} and "
        f"{fr(best_recall('IVF', W)['recall_at_10'])} at the same nprobe = {np_top}. The ceiling is therefore set by PQ "
        f"compression error, not by the partition: once the true neighbours are in the probed cells, quantised "
        f"distances still mis-rank them. Compared with plain IVF, IVF+PQ saves memory: its largest index is "
        f"{fmb(max(ivfpq_sizes[S]))} MB against {fmb(ivf_size[S])} MB for IVF on SIFT1M "
        f"({fx(ivf_size[S] / max(ivfpq_sizes[S]))} smaller), and {fmb(max(ivfpq_sizes[W]))} MB against "
        f"{fmb(ivf_size[W])} MB on Wikipedia ({fx(ivf_size[W] / max(ivfpq_sizes[W]))}). Re-ranking a short list of "
        f"IVF+PQ candidates with exact distances is the standard way to lift this ceiling (Section 4).")
    R.table("ivfpq", "PQ (ADC) against IVF+PQ at equal m. 'Match' is the fastest IVF+PQ configuration whose recall "
            "is at least PQ-ADC's; speed-up is PQ-ADC time divided by that time.",
            ["m", "PQ recall", "PQ ms", "IVF+PQ max recall (nprobe)", "IVF+PQ match ms (nprobe)", "Speed-up",
             "PQ MB", "IVF+PQ MB"],
            ivfpq_rows, widths=[1.0, 1.6, 1.4, 2.8, 2.8, 1.6, 1.5, 1.8])

    R.h2("3.8  Ablation: distance metric versus dataset in Annoy")
    mis = COSINE_GT["max_recall_loss_points"]
    by_k = dict(zip(map(int, abl_k), abl_diff))
    claim("Ablation: drop at large search_k within cosine/L2 mismatch; clear only at smallest",
          all(by_k[k] <= mis for k in sorted(by_k)[-2:]) and by_k[min(by_k)] > 1.5 * mis,
          f"mismatch {mis:.2f}; drops {', '.join(f'{k}: {d:.2f}' for k, d in sorted(by_k.items()))}")
    abl_rows = []
    for k, d, g in zip(abl_k, abl_diff, wiki_gap):
        abl_rows.append([fint(k), fr(ABLATION["euclidean"][k]["recall_at_10"]), fr(ABLATION["angular"][k]["recall_at_10"]),
                         f"{d:.1f}", fr(one("Annoy", W, k)["recall_at_10"]), f"{g:.1f}"])
    R.p(f"Annoy's recall is much lower on Wikipedia than on SIFT1M at small search_k "
        f"({fr(annoy_min[W]['recall_at_10'])} against {fr(annoy_min[S]['recall_at_10'])} at search_k = "
        f"{fint(min(annoy_mod.SEARCH_K_SWEEP))}), but the two runs differ in both data and metric (angular against "
        f"euclidean). To separate the two, we built a second SIFT1M index with the angular metric and scored both "
        f"against the same euclidean ground truth, so the data and dimensionality are fixed and only the metric "
        f"changes ({T('ablation')}, {F('ablation')}). Switching the metric "
        f"costs "
        f"{min(abl_diff):.1f}–{max(abl_diff):.1f} recall points, while the SIFT1M–Wikipedia gap is "
        f"{min(wiki_gap):.1f}–{max(wiki_gap):.1f} points and largest at small search_k. "
        f"Scoring the angular index against euclidean ground truth "
        f"slightly penalises it. Ranking SIFT1M by cosine instead changes {fint(COSINE_GT['differing_slots'])} of "
        f"the {fint(COSINE_GT['slots'])} true-neighbour slots ({COSINE_GT['max_recall_loss_points']:.1f} recall "
        f"points): base-vector norms are nearly constant ({COSINE_GT['base_norm']['mean']:.1f} ± "
        f"{COSINE_GT['base_norm']['std']:.1f}), but cosine and euclidean still disagree among near-ties. The "
        f"{min(abl_diff):.1f}–{max(abl_diff):.1f}-point drop is therefore an upper bound on the metric's effect. "
        f"At large search_k, where the angular index is close to exact, the drop is no larger than this mismatch, "
        f"so a clear effect remains only at search_k = {fint(min(annoy_mod.SEARCH_K_SWEEP))}. So the metric explains "
        f"only a small part of the gap. Most of it comes from the dataset. The Wikipedia vectors are three times "
        f"higher-dimensional, which may make random-projection splits less informative, but the datasets also differ "
        f"in distribution and size. Our ablation isolates the metric, not dimensionality, so attributing the "
        f"remainder to dimensionality is a hypothesis. The ablation was produced by a separate, earlier run "
        f"(committed before the single-thread settings were added), so we compare recall only, not its timings.")
    R.table("ablation", "Annoy recall@10 on SIFT1M with the euclidean and angular metrics (same data and ground "
            "truth), and on Wikipedia (angular, final run). Differences are in recall points (×100).",
            ["search_k", "SIFT1M eucl.", "SIFT1M ang.", "Metric Δ", "Wikipedia ang.", "SIFT1M−Wiki Δ"], abl_rows,
            widths=[2.0, 2.4, 2.4, 1.9, 2.6, 3.0])
    R.figure("ablation", os.path.join(FIG_DIR, "annoy_metric_ablation.png"),
             "Annoy recall@10 against search_k. The two SIFT1M curves (euclidean and angular metric on identical "
             "data) nearly overlap, while the Wikipedia curve lies well below them at small search_k, so most of the gap "
             "comes from the dataset rather than the metric.", width_cm=11.5)

    R.h2("3.9  Success and failure cases on Wikipedia")
    cs, ex = CASES["settings"], CASES["examples"]
    pq_rec = np.array(CASES["per_query"]["recall_ivfpq"])
    iv_rec = np.array(CASES["per_query"]["recall_ivf"])
    qt = CASES["s10_quartiles"]
    R.p(f"Mean recall hides how individual queries fare. We evaluated IVF (nprobe = {cs['ivf_nprobe']}) and "
        f"IVF+PQ (m = {cs['ivfpq_m']}, nprobe = {cs['ivfpq_nprobe']}) on every Wikipedia query "
        f"(`results/wiki_query_cases.py`; mean recall {fr(CASES['mean_recall_ivf'])} and "
        f"{fr(CASES['mean_recall_ivfpq'])}, as in the main run). For each query we also measured how tight its "
        f"neighbourhood is, as the cosine similarity of its true 10th neighbour (s10), and how many IVF cells its "
        f"true top-10 are spread over.")
    R.p(("Because IVF computes exact distances inside the cells it visits, its per-query recall equals the share "
         "of the true top-10 that lie in the probed cells; this holds for all "
         f"{fint(len(iv_rec))} queries. IVF fails only when true neighbours sit in cells it did not visit. "
         if CASES["ivf_recall_equals_coverage"] else
         "[CHECK: IVF recall did not equal probed-cell coverage for every query.] ")
        + f"{F('cases')}(a) shows the outcome: IVF is perfect on {int((iv_rec == 1).sum())} queries, whereas "
        f"IVF+PQ is perfect on only {int((pq_rec == 1).sum())}, because even when the right cells are visited, "
        f"compressed distances rarely rank all ten neighbours correctly. {F('cases')}(b) shows what makes a query "
        f"hard. In the quartile with the loosest neighbourhoods (s10 {qt[0]['s10_range'][0]:.2f}–"
        f"{qt[0]['s10_range'][1]:.2f}) the true neighbours spread over {qt[0]['mean_n_cells']:.1f} cells on "
        f"average and IVF recall is {fr(qt[0]['mean_recall_ivf'])}. In the tightest quartile (s10 "
        f"{qt[-1]['s10_range'][0]:.2f}–{qt[-1]['s10_range'][1]:.2f}) they span {qt[-1]['mean_n_cells']:.1f} "
        f"cells and recall is {fr(qt[-1]['mean_recall_ivf'])}. IVF+PQ follows the same trend, but its gap to IVF "
        f"widens for tight queries ({fr(qt[-1]['mean_recall_ivfpq'])} against {fr(qt[-1]['mean_recall_ivf'])}): "
        f"once coverage is nearly complete, compression error is what remains.")
    s_, f_ = ex["success"], ex["failure"]
    R.p(f"{T('cases')} gives one example of each, chosen by fixed rules (success: perfect recall for both methods "
        f"with the tightest neighbourhood; failure: lowest IVF+PQ recall). The success is a templated census "
        f"sentence whose true neighbours are near-identical sentences (s10 = {s_['s10']:.3f}) all in "
        f"{s_['n_cells']} cell, so both methods return them exactly. The failure is a long sentence combining "
        f"several topics; its closest match in the corpus is only loosely related (cosine {f_['s1']:.3f}), "
        f"and its true neighbours are scattered over {f_['n_cells']} cells, of which the probed cells hold "
        f"{f_['coverage'] * 100:.0f}% (IVF recall {f_['recall_ivf']:.1f}, IVF+PQ {f_['recall_ivfpq']:.1f}). "
        f"IVF+PQ keeps the top match but fills the rest with off-topic sentences. Hard queries for partition-based "
        f"methods tend to be queries with no close paraphrase in the corpus; probing more cells for such queries "
        f"(for example, adapting nprobe to the query's distance from its nearest centroids) is a natural "
        f"improvement.")

    def clip(text, n=115):
        return text if len(text) <= n else text[:n - 1].rstrip() + "…"

    case_rows = []
    for name, e in (("Success", s_), ("Failure", f_)):
        case_rows.append(f"{name} (query {e['query_index']}): recall IVF {e['recall_ivf']:.1f}, IVF+PQ "
                         f"{e['recall_ivfpq']:.1f}; s10 = {e['s10']:.3f}; true top-10 in {e['n_cells']} cell"
                         f"{'s' if e['n_cells'] != 1 else ''}")
        case_rows.append(["Query", "", clip(e["query"])])
        case_rows += [[f"True #{j + 1}", f"{t['cosine']:.3f}", clip(t["sentence"])] for j, t in enumerate(e["true_top"])]
        case_rows += [[f"IVF+PQ #{j + 1}", f"{t['cosine']:.3f}", clip(t["sentence"])] for j, t in enumerate(e["ivfpq_top"])]
    R.table("cases", "Success and failure example on Wikipedia: the query, its true top-3 neighbours and the "
            "top-3 returned by IVF+PQ, with cosine similarity to the query.",
            ["", "Cosine", "Sentence"], case_rows, align=["l", "r", "l"], widths=[2.0, 1.3, 12.6])
    R.figure("cases", os.path.join(FIG_DIR, "wiki_query_cases.png"),
             f"Per-query recall on Wikipedia for IVF (nprobe {cs['ivf_nprobe']}) and IVF+PQ (m {cs['ivfpq_m']}, "
             f"nprobe {cs['ivfpq_nprobe']}). (a) Most IVF failures are partial, and IVF+PQ is rarely perfect. "
             f"(b) Queries are grouped by the cosine similarity of their true 10th neighbour; recall rises steadily "
             f"as neighbourhoods tighten and the true neighbours concentrate in fewer cells.")

    R.h2("3.10  Discussion")
    R.p(f"**LSH is limited by its implementation as well as its algorithm.** On Wikipedia even the most selective "
        f"configuration re-ranks only {lsh_w_low['candidate_pct']:.1f}% of the base "
        f"({fint(lsh_w_low['avg_candidates'])} candidates) yet takes {fms(lsh_w_low['avg_query_time_ms'])} ms, no faster "
        f"than scanning all {fint(nb[W])} vectors ({fms(lin_ms[W])} ms; {RETIME_LSH}). That is about {ns_per_cand_w:.0f} ns per "
        f"candidate against {ns_per_vec_w:.0f} ns per vector for the BLAS-backed linear scan "
        f"(SIFT1M: {ns_per_cand_s:.0f} against {ns_per_vec_s:.0f} ns). Most of the per-query time goes to work "
        f"outside the distance computation: hashing, Python dictionary lookups, concatenating and de-duplicating "
        f"bucket lists, and gathering candidate rows. We did not profile these steps individually, so this attribution is "
        f"an inference from the per-candidate cost. A compiled or vectorised bucket store would shift the LSH curve "
        f"left, but LSH also needs far more candidates: on SIFT1M it re-ranks "
        f"{lsh_s['high']['candidate_pct']:.1f}% of the base for {fr(lsh_s['high']['recall_at_10'])}, whereas IVF "
        f"reaches {fr(iv[S]['recall_at_10'])} visiting {ivf_nprobe(iv[S]['param'])} of "
        f"{pqivf_mod.SIFT_CFG['nlist']} cells ({100 * ivf_nprobe(iv[S]['param']) / pqivf_mod.SIFT_CFG['nlist']:.1f}% "
        f"of the base if cells were equal-sized).")
    R.p(f"**Near-exact approximate search offers no speed-up over brute force.** At search_k = "
        f"{fint(max(annoy_mod.SEARCH_K_SWEEP))}, Annoy on Wikipedia reaches {fr(annoy_max[W]['recall_at_10'])} in "
        f"{fms(annoy_max[W]['avg_query_time_ms'])} ms against {fms(lin_ms[W])} ms for the exact linear scan, while "
        f"still not exact. {RETIME_ANNOY} Tree traversal and candidate handling cost more per vector than a dense matrix–vector "
        f"product, so the advantage of an ANN index disappears once it must inspect a large share of the data. "
        f"The same is true of every LSH configuration on Wikipedia.")
    R.p(f"**SDC is slower than ADC here.** Theory favours ADC on accuracy (Jégou et al., 2011), and our results agree, but SDC is also "
        f"slower at every m (for example {fms(next(r for r in pts('PQ-SDC', S) if pq_m(r['param']) == max(pqivf_mod.SIFT_CFG['pq_ms']))['avg_query_time_ms'])} "
        f"against {fms(pq_top[S]['avg_query_time_ms'])} ms at m = {max(pqivf_mod.SIFT_CFG['pq_ms'])} on SIFT1M). "
        f"Memory traffic does not explain this: although SDC's precomputed centroid-to-centroid table is "
        f"m × 256 × 256 floats, a given query reads only the one 256-entry row of each sub-quantiser selected by "
        f"its own code, the same working set as ADC's per-query table. The likelier cause is implementation: in "
        f"FAISS, the ADC scan over 8-bit codes runs a specialised kernel (`pq_scan_8bit`), whereas SDC uses a "
        f"generic lookup loop and must also encode the query first. We have not profiled this, so it remains a "
        f"hypothesis.")
    R.p(f"**IVF+PQ's frontier is stepped** because its Pareto-optimal configurations alternate between code sizes. "
        f"Walking up the SIFT1M frontier, the best configuration switches m in the order "
        f"{' → '.join(map(str, frontier_ms(S)))}, and on Wikipedia {' → '.join(map(str, frontier_ms(W)))}. Within "
        f"one m, raising nprobe gives a short rise that flattens at that m's ceiling, so the frontier consists of "
        f"pieces from successive m values.")
    R.p(f"**Linear scan scores {fr(lin[S]['recall_at_10'])} rather than 1 on SIFT1M.** This corresponds to "
        f"{misses_s} of the {fint(10 * nq[S])} ground-truth neighbour slots, and is likely due to distance ties "
        f"(SIFT components are integers, so several base vectors can be equidistant from a query) and float32 "
        f"rounding in the expanded distance, which can reorder near-ties relative to the provided ground truth. On "
        f"Wikipedia, whose ground truth we computed with the same arithmetic, linear scan scores "
        f"{fr(lin[W]['recall_at_10'])}.")

    hn_size = {ds: pts("HNSW", ds)[0]["index_size_mb"] for ds in DS}
    an_size = {ds: pts("Annoy", ds)[0]["index_size_mb"] for ds in DS}
    km_ivf_max = max(v for (ds, label), v in km_spread.items() if label.startswith("IVF, "))
    R.p(f"{T('strengths')} summarises the strengths and weaknesses of each method as observed in our experiments, "
        f"and the setting in which we would choose it.")
    R.table("strengths", "Strengths and weaknesses of each method in our experiments (values given as SIFT1M / "
            "Wikipedia where they differ).",
            ["Method", "Strengths", "Weaknesses", "Use when"],
            [["Linear scan",
              f"Exact (recall {fr(lin[S]['recall_at_10'])} / {fr(lin[W]['recall_at_10'])}); no build or training; "
              f"no parameters.",
              f"Cost grows with n·d: {fms(lin_ms[S])} / {fms(lin_ms[W])} ms per query; stores all raw vectors.",
              "Small collections, exact answers, ground truth."],
             ["LSH",
              "No training; simple; L, K and w tune the trade-off.",
              f"Needs {lsh_s['high']['candidate_pct']:.1f}% / {lsh_w['very_high']['candidate_pct']:.1f}% of the base "
              f"as candidates to reach about 0.9 recall; w must be tuned per dataset; {LSH_VS_LINEAR}.",
              "Not in this form; a compiled multi-probe variant might be competitive."],
             ["PQ (ADC / SDC)",
              f"Smallest index ({rng(pq_sizes[S] + pq_sizes[W], fmb)} MB); a single parameter, m.",
              f"Scans every code, so time grows with m; recall at most {fr(pq_top[S]['recall_at_10'])} / "
              f"{fr(pq_top[W]['recall_at_10'])}; SDC is less accurate and slower than ADC.",
              "Severe memory limits with modest recall needs."],
             ["IVF",
              f"Within about 2× of HNSW's query time on SIFT1M and tied on Wikipedia; recall up to {fr(best_recall('IVF', S)['recall_at_10'])} / "
              f"{fr(best_recall('IVF', W)['recall_at_10'])}; builds in {fs(ivf_build[S])} / {fs(ivf_build[W])} s; "
              f"insensitive to k-means seed (≤ {km_ivf_max:.1f} pts).",
              f"Stores raw vectors ({fmb(ivf_size[S])} / {fmb(ivf_size[W])} MB); misses neighbours that fall in "
              f"unprobed cells, so queries with loose neighbourhoods fail (Section 3.9).",
              "General-purpose default, especially when the index is rebuilt often."],
             ["IVF+PQ",
              f"Fast (from {fms(fastest_ivfpq)} ms) and compact ({rng(ivfpq_sizes[S] + ivfpq_sizes[W], fmb)} MB); "
              f"more accurate than PQ in {n_wins} of {len(pq_vs_ivfpq)} settings.",
              f"Recall ceiling {fr(ivfpq_best[S]['recall_at_10'])} / {fr(ivfpq_best[W]['recall_at_10'])} from "
              f"compression; two coupled parameters (m, nprobe).",
              "Memory-limited deployments; add exact re-ranking for high recall."],
             ["HNSW",
              f"Fastest to 0.9 recall on SIFT1M and tied on Wikipedia; recall up to "
              f"{fr(best_recall('HNSW', S)['recall_at_10'])} / {fr(best_recall('HNSW', W)['recall_at_10'])}; no "
              f"training.",
              f"Slowest build ({fs(hnsw_build[S])} / {fs(hnsw_build[W])} s single-threaded); index larger than the "
              f"raw data ({fmb(hn_size[S])} / {fmb(hn_size[W])} MB).",
              "Latency-critical search over a mostly static collection."],
             ["Annoy",
              f"Memory-mapped, shareable read-only index; recall up to {fr(annoy_max[S]['recall_at_10'])} / "
              f"{fr(annoy_max[W]['recall_at_10'])}.",
              f"Largest index ({fmb(an_size[S])} / {fmb(an_size[W])} MB); poor at its default search_k "
              f"({fr(ANNOY_DEFAULT[S]['recall_at_10'])} / {fr(ANNOY_DEFAULT[W]['recall_at_10'])}); {ANNOY_VS_LINEAR}.",
              "Read-only indexes shared across processes or loaded from disk."]],
            align=["l", "l", "l", "l"], widths=[2.0, 4.6, 5.6, 3.7])

    R.h2("3.11  Limitations")
    R.bullets([
        "The main results come from a single run per configuration on one machine (Apple M3), without confidence "
        "intervals, and absolute times are hardware-specific. The repeated timings below show how much "
        "those single runs can vary.",
        f"Wikipedia timings are noisy. We re-timed six Wikipedia configurations in {NUM[n_rt]} interleaved rounds "
        f"(`results/timing_controlled.py`), on mains power but with other applications open (load average "
        f"{min(rt_load):.1f}–{max(rt_load):.1f}). Each Wikipedia measurement is short (about "
        f"{np.mean(rt_short_s):.1f} s in total for IVF and HNSW), so brief background activity distorts it, which is "
        f"why those two varied most between rounds ({min(rt_spread[n] for n in rt_fast):.0f}–"
        f"{max(rt_spread[n] for n in rt_fast):.0f}%) and the configurations taking milliseconds per query varied "
        f"less ({min(rt_spread[n] for n in rt_slow):.0f}–{max(rt_spread[n] for n in rt_slow):.0f}%). The main run "
        f"was faster than every re-timing for {NUM[len(rt_main_faster)]} of the six configurations"
        + (f" (all but Annoy at search_k = {fint(max(annoy_mod.SEARCH_K_SWEEP))})"
           if set(rt) - set(rt_main_faster) == {"annoy_search_k100000"} else "")
        + f", so absolute Wikipedia times are indicative only and the within-round speed-ups are the more reliable "
        f"comparison. IVF at nprobe 16 was {rt_range('ivf_nprobe16')} faster than linear scan (median "
        f"{fx(np.median(rt_sp['ivf_nprobe16']))}) and HNSW at ef_search 64 {rt_range('hnsw_ef64')} (median "
        f"{fx(np.median(rt_sp['hnsw_ef64']))}), against {fx(lin_ms[W] / iv[W]['avg_query_time_ms'])} and "
        f"{fx(lin_ms[W] / hn[W]['avg_query_time_ms'])} in the main run; IVF was the faster of the two in "
        f"{NUM[int((rt['ivf_nprobe16'] < rt['hnsw_ef64']).sum())]} of {NUM[n_rt]} rounds, so we treat them as tied. "
        f"LSH \"low\" ({rt_range('lsh_low', 2)}, median {np.median(rt_sp['lsh_low']):.2f}×) and Annoy at search_k = "
        f"{fint(max(annoy_mod.SEARCH_K_SWEEP))} ({rt_range('annoy_search_k100000', 2)}, median "
        f"{np.median(rt_sp['annoy_search_k100000']):.2f}×) stayed level with linear scan in every round. SIFT1M, "
        f"with {fint(nq[S])} queries per measurement, varied less between rounds "
        f"({min(rts_spread.values()):.0f}–{max(rts_spread.values()):.0f}%) but showed the same pattern. Re-timed "
        f"HNSW (ef_search 32) and IVF (nprobe 16) ran at {rt_range('hnsw_ef32', 0, rts_sp)} (median "
        f"{fx(rt_med[S][0])}) and {rt_range('ivf_nprobe16', 0, rts_sp)} (median {fx(rt_med[S][1])}) the speed of "
        f"linear scan, against {fx(lin_ms[S] / hn[S]['avg_query_time_ms'])} and "
        f"{fx(lin_ms[S] / iv[S]['avg_query_time_ms'])} in the main run, while Annoy at search_k = 5,000 ran at "
        f"{rt_range('annoy_search_k5000', 0, rts_sp)} (median {fx(np.median(rts_sp['annoy_search_k5000']))}; main "
        f"run {fx(lin_ms[S] / one('Annoy', S, '5000')['avg_query_time_ms'])}). HNSW, IVF and Annoy kept that order "
        f"in every round, so our SIFT1M conclusions stand, but the main-run speed-ups for HNSW and IVF are at the "
        f"optimistic end.",
        "Single-threaded, one-query-at-a-time timing reflects latency, not throughput. Batched or multi-threaded "
        "search would favour the FAISS and HNSW implementations, which parallelise across queries.",
        "HNSW (random level assignment) and Annoy (random splits) are randomised builds. A rebuild can shift recall "
        f"slightly: Annoy's SIFT1M recall at search_k = {fint(min(annoy_mod.SEARCH_K_SWEEP))} differed by "
        f"{abs(ABLATION['euclidean'][str(min(annoy_mod.SEARCH_K_SWEEP))]['recall_at_10'] - annoy_min[S]['recall_at_10']):.4f} "
        f"between two builds ({ABLATION['euclidean'][str(min(annoy_mod.SEARCH_K_SWEEP))]['recall_at_10']:.4f} in the "
        f"ablation run, {annoy_min[S]['recall_at_10']:.4f} in the final run).",
        "LSH is a pure-Python/NumPy implementation, while the other ANN methods use optimised C++ libraries, so "
        "its query times overstate the algorithm's inherent cost.",
        "The applied dataset uses a single embedding model and 1,000 queries. Other models and corpora may change "
        "the relative standing of the methods, particularly for tree- and hash-based methods.",
        "Parameter grids are coarse (for example three or four LSH configurations and four Annoy search_k values), "
        "so the frontiers are piecewise approximations.",
        f"Query times include Python function-call overhead. The fastest configurations (down to "
        f"{fms(fastest_ivfpq)} ms for IVF+PQ) are therefore close to the measurement floor, and small differences "
        f"among them should not be over-interpreted; comparisons at 0.1 ms and above are not materially affected.",
    ])

    # ------------------------------------------------------------------ Conclusion
    R.h1("4  Conclusion")
    R.p(f"We compared six ANN methods and an exact baseline under identical single-threaded conditions on SIFT1M and "
        f"on Wikipedia sentence embeddings. Graph and partition methods dominate the recall/latency trade-off. On "
        f"SIFT1M, {hnsw_ivf(S)}; on Wikipedia, {hnsw_ivf(W)} {RT_NOTE}. IVF also builds two orders of magnitude faster than "
        f"HNSW. Compression trades "
        f"recall for memory: PQ-based indexes are one to two orders of magnitude smaller but plateau below 0.9, "
        f"with IVF+PQ the better of the two in speed and, in {n_wins} of {len(pq_vs_ivfpq)} settings, in recall. Our LSH implementation is not competitive, and on Wikipedia it "
        f"is no faster than brute force at any setting.")
    R.p("**Practical recommendations.**").paragraph_format.keep_with_next = True
    R.bullets([
        "**Lowest latency at high recall, memory available:** HNSW, with ef_search as the operating-point knob. "
        "Choose IVF instead when the index must be rebuilt often, since its build is far cheaper and its query "
        "times are within about 2× of HNSW's (tied on Wikipedia).",
        "**Tight memory budget:** IVF+PQ, accepting a recall ceiling of roughly "
        f"{fr(ivfpq_best[S]['recall_at_10'])}–{fr(ivfpq_best[W]['recall_at_10'])} on these datasets, or adding exact "
        "re-ranking of a short candidate list to recover recall.",
        "**Read-only index shared across processes or loaded from disk:** Annoy, using search_k well above its "
        "default; on our data it needs the largest index and is slower than HNSW and IVF at equal recall.",
        f"**Exact answers or small collections:** linear scan, which at these sizes takes "
        f"{fms(lin_ms[S])}–{fms(lin_ms[W])} ms per query on one core and needs no index.",
        "**Avoid** basic LSH with pure-Python buckets for dense learned embeddings.",
    ])
    R.p("**Future work.** Multi-probe and vectorised LSH (compiled bucket storage, probing neighbouring buckets) to "
        "test how much of LSH's gap is implementation; OPQ (optimised rotations before PQ) and IVF+PQ with exact "
        "re-ranking to lift the compression ceiling; GPU and multi-threaded FAISS indexes and batched queries to "
        "measure throughput; repeated runs with confidence intervals; and larger scales (e.g. SIFT1B or Deep1B) and "
        "further embedding models.")

    # ------------------------------------------------------------------ References
    R.h1("References")
    # APA 7th edition: alphabetical by first author, hanging indent, DOIs
    # (each checked against Crossref; Sentence-BERT pages per the ACL Anthology).
    refs = [
        "Bernhardsson, E. (2013). *Annoy: Approximate nearest neighbors oh yeah* [Computer software]. Spotify. "
        "Retrieved October 2, 2026, from https://github.com/spotify/annoy",
        "Datar, M., Immorlica, N., Indyk, P., & Mirrokni, V. S. (2004). Locality-sensitive hashing scheme based on "
        "p-stable distributions. In *Proceedings of the twentieth annual Symposium on Computational Geometry* "
        "(pp. 253–262). Association for Computing Machinery. https://doi.org/10.1145/997817.997857",
        "Indyk, P., & Motwani, R. (1998). Approximate nearest neighbors: Towards removing the curse of "
        "dimensionality. In *Proceedings of the thirtieth annual ACM Symposium on Theory of Computing* "
        "(pp. 604–613). Association for Computing Machinery. https://doi.org/10.1145/276698.276876",
        "Jégou, H., Douze, M., & Schmid, C. (2011). Product quantization for nearest neighbor search. *IEEE "
        "Transactions on Pattern Analysis and Machine Intelligence*, *33*(1), 117–128. "
        "https://doi.org/10.1109/TPAMI.2010.57",
        "Johnson, J., Douze, M., & Jégou, H. (2021). Billion-scale similarity search with GPUs. *IEEE Transactions "
        "on Big Data*, *7*(3), 535–547. https://doi.org/10.1109/TBDATA.2019.2921572",
        "Lowe, D. G. (2004). Distinctive image features from scale-invariant keypoints. *International Journal of "
        "Computer Vision*, *60*(2), 91–110. https://doi.org/10.1023/B:VISI.0000029664.99615.94",
        "Malkov, Y. A., & Yashunin, D. A. (2020). Efficient and robust approximate nearest neighbor search using "
        "hierarchical navigable small world graphs. *IEEE Transactions on Pattern Analysis and Machine "
        "Intelligence*, *42*(4), 824–836. https://doi.org/10.1109/TPAMI.2018.2889473",
        "Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence embeddings using Siamese BERT-networks. In "
        "*Proceedings of the 2019 Conference on Empirical Methods in Natural Language Processing and the 9th "
        "International Joint Conference on Natural Language Processing (EMNLP-IJCNLP)* (pp. 3982–3992). "
        "Association for Computational Linguistics. https://doi.org/10.18653/v1/D19-1410",
        "Sentence Transformers. (n.d.). *wikipedia-en-sentences* [Data set]. Hugging Face. Retrieved October 2, "
        "2026, from https://huggingface.co/datasets/sentence-transformers/wikipedia-en-sentences",
    ]
    for ref in refs:
        par = R.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(1.27)
        par.paragraph_format.first_line_indent = Cm(-1.27)
        par.paragraph_format.space_after = Pt(4)
        add_runs(par, ref, 10)

    R.save(path_docx)
    return R


def collect_checks(path_docx):
    doc = Document(path_docx)
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        texts += [c.text for row in t.rows for c in row.cells]
    out = []
    for t in texts:
        out += re.findall(r"\[CHECK[^\]]*\]", t)
    return list(dict.fromkeys(out))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", action="store_true", help="convert to PDF with Microsoft Word")
    args = ap.parse_args()
    os.makedirs(FIG_DIR, exist_ok=True)
    for ds in DS:
        fig_recall_latency(ds, os.path.join(FIG_DIR, f"recall_vs_latency_{ds}.png"))
    fig_index_size(os.path.join(FIG_DIR, "index_size.png"))
    fig_ablation(os.path.join(FIG_DIR, "annoy_metric_ablation.png"))
    fig_cases(os.path.join(FIG_DIR, "wiki_query_cases.png"))

    check_claims()
    docx_path = os.path.join(REPORT_DIR, "group_XX_report.docx")
    build(docx_path)
    print("saved", docx_path)

    print("\nClaim checks:")
    for text, ok, detail in CLAIMS:
        print(f"  [{'ok' if ok else 'NOT SUPPORTED'}] {text}: {detail}")
    print("\n[CHECK] markers:")
    for c in collect_checks(docx_path):
        print("  ", c)

    if args.pdf:
        from docx2pdf import convert
        from pypdf import PdfReader
        pdf_path = docx_path[:-5] + ".pdf"
        convert(docx_path, pdf_path)
        print("\nsaved", pdf_path, "|", len(PdfReader(pdf_path).pages), "pages")


if __name__ == "__main__":
    main()
